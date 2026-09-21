from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def paragraph_info(paragraph):
    return {
        "style": paragraph.style.name if paragraph.style else None,
        "text": paragraph.text,
        "alignment": str(paragraph.alignment),
    }


def cell_info(cell):
    return {
        "text": cell.text,
        "paragraphs": [paragraph_info(p) for p in cell.paragraphs],
        "nested_tables": [table_info(t) for t in cell.tables],
    }


def table_info(table):
    return {
        "alignment": str(table.alignment),
        "autofit": table.autofit,
        "column_widths": [str(cell.width) for cell in table.rows[0].cells],
        "rows": [
            [cell_info(cell) for cell in row.cells]
            for row in table.rows
        ]
    }


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    path = Path(sys.argv[1])
    document = Document(path)
    result = {
        "path": str(path),
        "paragraphs": [paragraph_info(p) for p in document.paragraphs],
        "tables": [table_info(t) for t in document.tables],
        "sections": [],
        "inline_shapes": len(document.inline_shapes),
        "core_properties": {
            "title": document.core_properties.title,
            "subject": document.core_properties.subject,
            "author": document.core_properties.author,
            "last_modified_by": document.core_properties.last_modified_by,
        },
    }
    for section in document.sections:
        result["sections"].append({
            "header": [paragraph_info(p) for p in section.header.paragraphs],
            "footer": [paragraph_info(p) for p in section.footer.paragraphs],
            "page_width": section.page_width,
            "page_height": section.page_height,
            "top_margin": section.top_margin,
            "bottom_margin": section.bottom_margin,
            "left_margin": section.left_margin,
            "right_margin": section.right_margin,
        })
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
