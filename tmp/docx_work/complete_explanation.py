from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE = Path(r"D:\Downloads\BM.ĐT.19.21 BẢN GIẢI TRÌNH SỬA CHỮA DAKLTN.docx")
OUTPUT = (
    Path(r"C:\Program1\Code\Unity\GraduationProject\output\documents")
    / "BM.ĐT.19.21 BẢN GIẢI TRÌNH SỬA CHỮA DAKLTN - HOÀN THIỆN.docx"
)


CORRECTIONS = [
    (
        "Độ tin cậy thống kê và yêu cầu lặp nhiều seed",
        "68-97; 124-128",
        (
            "Đã làm rõ toàn bộ 15 tổ hợp thực nghiệm hiện mới được chạy với một seed, "
            "vì vậy các kết quả chỉ có giá trị mô tả và thăm dò. Đã loại bỏ cách xem "
            "các điểm TensorBoard liên tiếp trong cùng một lần chạy là các mẫu độc lập, "
            "cũng như các kết luận kiểm định chưa đủ cơ sở. Đồ án bổ sung giao thức lặp "
            "tối thiểu 5 seed, sử dụng cùng danh sách seed cho các thuật toán, báo cáo "
            "trung bình, độ lệch chuẩn giữa seed, khoảng tin cậy 95% và toàn bộ đường "
            "cong. Việc chạy bổ sung nhiều seed được nêu là hướng thực nghiệm tiếp theo, "
            "không trình bày như dữ liệu đã có."
        ),
        "68-97; 127-133",
    ),
    (
        "Tăng độ phân giải phép đo trên Football Table",
        "73-80; 92-94; 124-128",
        (
            "Đã bổ sung phân tích: ELO mỗi lần self-play dùng kho đối thủ riêng, không thể "
            "dùng trực tiếp để xếp hạng. Hướng đánh giá gồm: tăng ngân sách trên 1,6 triệu "
            "bước; lưu snapshot; tổ chức vòng tròn trước cùng bot luật và checkpoint cố "
            "định; đổi phía; báo cáo tỷ lệ thắng, hòa, chênh lệch bàn và khoảng tin cậy."
        ),
        "73-80; 92-94; 127-132",
    ),
    (
        "Phân tích SAC và chênh lệch giữa các thuật toán trên Capture The Flag",
        "84-90; 95-97; 127-128",
        (
            "Đã điều chỉnh cách diễn đạt từ nhận định khái quát về SAC sang mô tả chính "
            "xác rằng lần chạy SAC được ghi nhận không hội tụ. Đồ án bổ sung các giả "
            "thuyết cần kiểm chứng liên quan đến curriculum, replay buffer, phần thưởng "
            "nhóm, entropy và tần suất cập nhật; đồng thời đề xuất các thí nghiệm ablation "
            "gồm bỏ curriculum, xóa hoặc chia replay buffer khi chuyển bậc, thay đổi "
            "entropy/tần suất cập nhật và chạy cùng seed, cùng ngân sách. Chênh lệch nhỏ "
            "giữa các thuật toán đạt kết quả cao được xác định là chưa đủ để xếp hạng nếu "
            "chưa có đánh giá nhiều seed trên tập episode cố định."
        ),
        "84-90; 95-97; 130-132",
    ),
    (
        "Trình bày riêng DQN trên Football Table",
        "73-80; 94-96; 108; 127-128",
        (
            "Đã tách DQN khỏi nhóm bốn thuật toán dùng self-play. DQN không còn xuất "
            "hiện trong biểu đồ/xếp hạng ELO hay cột so sánh tổng hợp của Football Table; "
            "kết quả được chuyển sang bảng tham chiếu riêng, ghi rõ dùng scene hành động "
            "rời rạc, không self-play và không có ELO. Trong ứng dụng trình diễn, cột DQN "
            "được làm mờ, gắn dấu sao và nhãn 'Riêng', đồng thời bị loại khỏi phép sắp "
            "hạng với các thuật toán còn lại."
        ),
        "74-76; 95-97; 108; 130-131",
    ),
    (
        "Kiểm thử ngoài Unity Editor và kiểm thử tự động",
        "121-123; 124-128",
        (
            "Đã nêu rõ phạm vi hiện tại mới là kiểm thử thủ công trong Unity Editor trên "
            "Windows, chưa thay thế kiểm thử bản build hoặc kiểm thử tự động trên máy "
            "sạch. Đồ án bổ sung checklist và lộ trình kiểm thử bằng Unity Test Framework/"
            "PlayMode: cài package vào dự án sạch, import sample, build standalone, nạp "
            "model ONNX liên tục và rời rạc, xử lý model sai đặc tả, kiểm tra dữ liệu "
            "TensorBoard/JSON và các luồng nạp scene, quay lại menu."
        ),
        "120-125; 127-133",
    ),
    (
        "Hướng dẫn sử dụng package, phần cứng, thời gian và giới hạn tái sử dụng",
        "115-117",
        (
            "Đã bổ sung các tiểu mục độc lập trong Mục 5.7: quy trình cài đặt, sử dụng UPM "
            "và nạp ONNX; yêu cầu Unity, ML-Agents, Inference Engine, Python/PyTorch; cấu "
            "hình Windows 11, Core i5-14600KF, RTX 3050, 32 GB RAM; thời gian huấn luyện "
            "dự kiến cho 15 cấu hình; giới hạn về một seed, DQN không self-play, trainer "
            "plugin, art Asset Store, thay đổi đặc tả ONNX và nền tảng."
        ),
        "116-121",
    ),
    (
        "Rà soát thuật ngữ, số liệu hình/bảng và cỡ chữ",
        "Toàn văn",
        (
            "Đã thống nhất cách ghi PPO, SAC, MA-POCA, MAPPO, DQN và tên ba môi trường "
            "Football Table, Capture The Flag, Cross The Road; đối chiếu lại số liệu giữa "
            "phần mô tả, hình, bảng và phụ lục; sửa các tham chiếu chương viết cứng; tăng "
            "cỡ chữ nhãn số trên một số biểu đồ từ mức quá nhỏ lên mức dễ đọc hơn. Bản "
            "LaTeX sau chỉnh sửa đã được biên dịch lại, không còn tham chiếu/citation chưa "
            "định nghĩa và không có nội dung tràn lề."
        ),
        "Toàn văn; 68-125",
    ),
]


def set_run_font(run, size=11, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def set_cell_margin(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color="000000", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def set_table_fixed_layout(table):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = tr_pr.find(qn("w:tblHeader"))
    if marker is None:
        marker = OxmlElement("w:tblHeader")
        tr_pr.append(marker)
    marker.set(qn("w:val"), "true")


def allow_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is not None:
        tr_pr.remove(cant_split)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    cant_split.set(qn("w:val"), "true")


def clear_row_height(row):
    tr_pr = row._tr.get_or_add_trPr()
    for height in list(tr_pr.findall(qn("w:trHeight"))):
        tr_pr.remove(height)


def set_cell_text(cell, text, *, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, size=10.5, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margin(cell)
    set_cell_borders(cell)


def set_column_widths(table, widths):
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = width
            tc_pr = row.cells[index]._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width / 635)))
            tc_w.set(qn("w:type"), "dxa")


def replace_paragraph_text(paragraph, text):
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run, size=13)


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)

    for paragraph in document.paragraphs:
        if "Đã bảo vệ đồ án/khốa luận" in paragraph.text:
            replace_paragraph_text(
                paragraph,
                "Đã bảo vệ đồ án/khóa luận tốt nghiệp ngày 03 tháng 09 năm 2026",
            )
        elif paragraph.text.startswith("Hà Nội, ngày"):
            replace_paragraph_text(paragraph, "Hà Nội, ngày 19 tháng 09 năm 2026")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in paragraph.runs:
                run.italic = True

    correction_table = document.tables[1]
    set_table_fixed_layout(correction_table)
    set_repeat_header(correction_table.rows[0])

    while len(correction_table.rows) < len(CORRECTIONS) + 1:
        correction_table.add_row()

    widths = [Inches(2.187), Inches(0.79), Inches(2.453), Inches(1.063)]
    set_column_widths(correction_table, widths)

    for header_cell in correction_table.rows[0].cells:
        clear_row_height(correction_table.rows[0])
        header_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margin(header_cell, top=110, start=80, bottom=110, end=80)
        set_cell_borders(header_cell)

    for item_index, (row, values) in enumerate(zip(correction_table.rows[1:], CORRECTIONS)):
        clear_row_height(row)
        if item_index == 5:
            prevent_row_split(row)
        else:
            allow_row_split(row)
        set_cell_text(row.cells[0], values[0], alignment=WD_ALIGN_PARAGRAPH.LEFT, size=10.5, bold=True)
        set_cell_text(row.cells[1], values[1], alignment=WD_ALIGN_PARAGRAPH.CENTER, size=10)
        set_cell_text(row.cells[2], values[2], alignment=WD_ALIGN_PARAGRAPH.LEFT, size=10.5)
        set_cell_text(row.cells[3], values[3], alignment=WD_ALIGN_PARAGRAPH.CENTER, size=10)

    position_table = document.tables[2]
    set_table_fixed_layout(position_table)
    set_repeat_header(position_table.rows[0])
    set_column_widths(position_table, [Inches(2.61), Inches(2.97), Inches(0.774)])
    for row in position_table.rows:
        clear_row_height(row)
        for cell in row.cells:
            set_cell_borders(cell)
    set_cell_text(position_table.rows[1].cells[0], "Không có", alignment=WD_ALIGN_PARAGRAPH.CENTER, size=11, bold=True)
    set_cell_text(
        position_table.rows[1].cells[1],
        "Tác giả đã tiếp thu và chỉnh sửa toàn bộ các góp ý của Hội đồng.",
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        size=11,
    )
    set_cell_text(position_table.rows[1].cells[2], "-", alignment=WD_ALIGN_PARAGRAPH.CENTER, size=11)

    for paragraph in document.paragraphs:
        if paragraph.text.startswith("2. Tác giả xin được giữ quan điểm"):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(3)
        if paragraph.text.startswith("Trên đây là Bản giải trình"):
            paragraph.paragraph_format.keep_with_next = True

    document.core_properties.title = "Bản giải trình sửa chữa đồ án/khóa luận tốt nghiệp"
    document.core_properties.subject = "Giải trình các nội dung chỉnh sửa theo góp ý của Hội đồng"
    document.core_properties.last_modified_by = "Văn Thị Minh Huyền"
    document.save(OUTPUT)
    print(str(OUTPUT))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
