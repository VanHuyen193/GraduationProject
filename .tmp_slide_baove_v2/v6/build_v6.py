# -*- coding: utf-8 -*-
"""Bản v6 của bộ slide bảo vệ: slide gọn hết mức, mọi diễn giải dồn vào kịch bản nói.

Khác v5:
  · bỏ slide case study Capture The Flag;
  · Football Table chỉ còn biểu đồ ELO, bỏ biểu đồ phần thưởng;
  · ba biểu đồ đường thay bằng ảnh chụp thẳng từ TensorBoard (xem tb_shot.py);
  · biểu đồ cột tổng hợp thay bằng bảng;
  · lời giải thích số liệu, khoảng ân hạn, ngưỡng phần thưởng… chuyển hết xuống
    Speaker Notes.

Chạy: $env:PYTHONUTF8="1"; C:\\Users\\Admin\\miniconda3\\python.exe build_v6.py
"""
import os
import sys

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from deckkit import *  # noqa

ROOT = os.path.dirname(os.path.abspath(__file__))
SRC = r"C:\Users\Admin\Downloads\DoAn\Slide_BaoVe_DoAnTotNghiep_v4.pptx"
OUT = os.path.join(ROOT, "Slide_BaoVe_DoAnTotNghiep_v6.pptx")
IMG = r"C:\Users\Admin\Documents\GitHub\GraduationProject\Đồ_án_tốt_nghiệp\image\app"
TB = os.path.join(ROOT, "shots")

RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
RED = RGBColor(0xC0, 0x39, 0x2B)
GREEN = RGBColor(0x27, 0x92, 0x4F)
SKY = RGBColor(0xCF, 0xD8, 0xEC)
AMBER = RGBColor(0xF2, 0x9F, 0x5B)

# --------------------------------------------------------------------- khởi tạo
prs = Presentation(SRC)
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.get(RID))
    sldIdLst.remove(sldId)

master = prs.slide_masters[0]
LAY = {l.name: l for l in master.slide_layouts}
LAY_TITLE = LAY["Title Slide"]
LAY_BODY = LAY["1_Title and Content"]

OVERFLOW = []


def new_slide_t(t, tsz=26, tw=9.90):
    s = prs.slides.add_slide(LAY_BODY)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    w = text_w(t, tsz, bold=True)
    while w > tw and tsz > 19.5:
        tsz -= 0.5
        w = text_w(t, tsz, bold=True)
    title(s, t, sz=tsz, w=tw)
    return s


def blank_title_slide():
    s = prs.slides.add_slide(LAY_TITLE)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    return s


def check(name, box_h, need_h):
    if need_h > box_h + 0.005:
        OVERFLOW.append(f"{name}: cần {need_h:.2f}in > khung {box_h:.2f}in")


def headline(slide, x, y, w, h, head, body=None, hsz=15.5, bsz=13):
    """Thẻ chỉ có tiêu đề, kèm tối đa một dòng phụ rất ngắn."""
    card(slide, x, y, w, h)
    paras = [{"runs": [(head, {"sz": hsz, "b": True, "c": NAVY})], "line": 1.04,
              "space_after": 4 if body else 0}]
    if body:
        paras.append({"runs": [(body, {"sz": bsz, "c": MUTED})], "line": 1.06})
    textbox(slide, x + 0.24, y, w - 0.48, h, paras, anchor="c")


def tb_plate(slide, x, y, w, name):
    """Đặt ảnh chụp TensorBoard, giữ đúng tỉ lệ gốc."""
    from PIL import Image
    path = os.path.join(TB, name)
    iw, ih = Image.open(path).size
    h = w * ih / iw
    round_rect(slide, x - 0.06, y - 0.06, w + 0.12, h + 0.12, fill=WHITE, adj=0.03)
    slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return h


def note_line(slide, x, y, w, h, n, head, body):
    card(slide, x, y, w, h)
    oval(slide, x + 0.20, y + 0.20, 0.30, fill=NAVY, label=str(n), sz=11)
    textbox(slide, x + 0.60, y + 0.14, w - 0.84, h - 0.28, [
        {"runs": [(head, {"sz": 13.5, "b": True, "c": NAVY})], "line": 1.02,
         "space_after": 3},
        {"runs": [(body, {"sz": 12.5, "c": INK})], "line": 1.06}], anchor="c")


# ============================================================== 1. TRANG BÌA ===
s = blank_title_slide()
simple(s, 1.00, 2.46, 11.33, 0.32,
       "ĐỒ ÁN TỐT NGHIỆP ĐẠI HỌC  •  NGÀNH KHOA HỌC MÁY TÍNH",
       sz=13, b=True, c=AMBER, align="c", wrap=False)
textbox(s, 1.00, 2.84, 11.33, 1.16, [
    {"runs": [("HỌC TĂNG CƯỜNG ĐỂ TỐI ƯU HÓA KHẢ NĂNG", {"sz": 30, "b": True, "c": WHITE})],
     "align": "c", "line": 1.12},
    {"runs": [("TƯƠNG TÁC CỦA AI TRONG GAME", {"sz": 30, "b": True, "c": WHITE})],
     "align": "c", "line": 1.12},
], anchor="t")

_env = [("env_football.png", "Football Table"),
        ("env_puzzle.png", "Capture The Flag"),
        ("env_crossroad.png", "Cross The Road")]
_w, _gap = 2.66, 0.30
_x0 = (SLIDE_W - (3 * _w + 2 * _gap)) / 2
for i, (fn, lab) in enumerate(_env):
    x = _x0 + i * (_w + _gap)
    picture_plate(s, x, 4.34, _w, 1.62, os.path.join(IMG, fn), pad=0.08)
    simple(s, x, 6.00, _w, 0.26, lab, sz=11, b=True, c=SKY, align="c", wrap=False)

textbox(s, 1.00, 6.36, 11.33, 0.70, [
    {"runs": [("Sinh viên thực hiện: ", {"sz": 14, "c": SKY}),
              ("Văn Thị Minh Huyền", {"sz": 14, "b": True, "c": WHITE}),
              ("   ·   MSSV: ", {"sz": 14, "c": SKY}),
              ("22010329", {"sz": 14, "b": True, "c": WHITE}),
              ("   ·   Khóa: ", {"sz": 14, "c": SKY}),
              ("K16", {"sz": 14, "b": True, "c": WHITE})],
     "align": "c", "line": 1.15, "space_after": 3},
    {"runs": [("Giảng viên hướng dẫn: ", {"sz": 14, "c": SKY}),
              ("ThS. Nguyễn Văn Sơn", {"sz": 14, "b": True, "c": WHITE})],
     "align": "c", "line": 1.15},
], anchor="t")
notes(s, """Kính chào quý thầy cô trong hội đồng. Em là Văn Thị Minh Huyền, sinh viên lớp K16
ngành Khoa học máy tính, mã sinh viên 22010329. Đề tài đồ án của em là "Học tăng cường để tối ưu
hóa khả năng tương tác của AI trong game", do thầy Nguyễn Văn Sơn hướng dẫn.

Bài trình bày của em đi theo bốn phần. Phần mở đầu nói về một quan sát đã định hình toàn bộ đồ án:
trong học tăng cường cho game, phần khó không nằm ở thuật toán mà nằm ở khâu đặc tả môi trường.
Phần tiếp theo trình bày ba môi trường em đã thiết kế và nguyên tắc thiết kế hàm phần thưởng.
Phần ba là cách em kiểm chứng ba môi trường đó bằng mười lăm lần huấn luyện. Phần cuối là kết quả
và luận điểm rút ra.

Ba ảnh phía dưới là ba môi trường em đã dựng trong Unity: Football Table là bàn bi lắc đối kháng,
Capture The Flag là câu đố hợp tác hai tác nhân, và Cross The Road là bài toán điều hướng qua
đường. Em xin phép bắt đầu.""")

# ====================================================== 2. BỐI CẢNH & VẤN ĐỀ ===
s = new_slide_t("Bối cảnh — từ logic viết tay tới học tăng cường")

L_X, L_W = 0.55, 6.30
_cards = [("Hiện nay NPC chạy bằng logic viết tay",
           "máy trạng thái, cây hành vi, luật if – else"),
          ("Được việc, nhưng hành vi cứng nhắc",
           "NPC lặp đúng kịch bản đã lường trước; người chơi đoán ra rồi khai thác"),
          ("Học tăng cường sinh hành vi từ tương tác",
           "đổi lại, phần khó chuyển sang khâu đặc tả môi trường")]
for i, (h, b) in enumerate(_cards):
    headline(s, L_X, 1.30 + i * 1.34, L_W, 1.18, h, b, hsz=16.5, bsz=13.5)

for i, (num, lab) in enumerate([("30+", "trang đặc tả ba môi trường"),
                                ("3", "trang cấu hình năm thuật toán")]):
    stat_tile(s, L_X + i * 3.24, 5.44, 3.06, 1.30, num, lab, nsz=34, lsz=12.5)

R_X, R_W = 7.15, 5.65
card(s, R_X, 1.30, R_W, 5.44)
simple(s, R_X + 0.24, 1.50, R_W - 0.48, 0.32, "VÒNG LẶP HỌC TĂNG CƯỜNG",
       sz=12.5, b=True, c=ORANGE, align="c", wrap=False)

_bx, _bw, _bh = R_X + 1.05, 3.55, 0.74
_boxes = [("TÁC NHÂN", "mạng chính sách π(a|s)", 2.10),
          ("MÔI TRƯỜNG UNITY", "cơ chế game · vật lý · va chạm", 3.40),
          ("HÀM PHẦN THƯỞNG", "tín hiệu duy nhất định nghĩa mục tiêu", 4.70)]
for i, (h, sub, by) in enumerate(_boxes):
    fill = NAVY if i == 2 else WHITE
    ln = None if i == 2 else NAVY
    round_rect(s, _bx, by, _bw, _bh, fill=fill, line=ln, lw=1.0, adj=0.12)
    textbox(s, _bx + 0.10, by + 0.06, _bw - 0.20, _bh - 0.12, [
        {"runs": [(h, {"sz": 13, "b": True, "c": WHITE if i == 2 else NAVY})],
         "align": "c", "line": 1.0, "space_after": 2},
        {"runs": [(sub, {"sz": 10.5, "c": ICE if i == 2 else MUTED})],
         "align": "c", "line": 1.0}], anchor="c")
    if i < 2:
        arrow(s, _bx + _bw / 2 - 0.11, by + _bh + 0.10, 0.22, 0.42, MSO_SHAPE.DOWN_ARROW)
simple(s, _bx + _bw / 2 + 0.18, 2.94, 1.60, 0.22, "hành động aₜ", sz=10, i=True,
       c=MUTED, wrap=False)
simple(s, _bx + _bw / 2 + 0.18, 4.24, 1.90, 0.22, "trạng thái sₜ₊₁", sz=10, i=True,
       c=MUTED, wrap=False)

_rail = RGBColor(0xB4, 0xBE, 0xD6)
rect(s, R_X + 0.42, 2.40, 0.05, 2.64, fill=_rail)
rect(s, R_X + 0.42, 4.99, 0.63, 0.05, fill=_rail)
rect(s, R_X + 0.42, 2.40, 0.45, 0.05, fill=_rail)
arrow(s, R_X + 0.85, 2.28, 0.24, 0.28, MSO_SHAPE.RIGHT_ARROW)
simple(s, R_X + 0.24, 5.62, R_W - 0.48, 0.60,
       "quan sát sₜ₊₁ và phần thưởng rₜ₊₁ quay lại tác nhân",
       sz=11.5, i=True, c=MUTED, align="c", wrap=False)
simple(s, R_X + 0.24, 6.10, R_W - 0.48, 0.56,
       "Người phát triển vẫn phải đặc tả cả ba ô — và đặc tả sai thì không có báo lỗi nào.",
       sz=12.5, b=True, c=NAVY, align="c", line=1.08)
caption(s, 0.55, 6.80, 12.25,
        "Nguồn: D. Isla, GDC 2005 (cây hành vi trong Halo 2)  ·  Yannakakis & Togelius, "
        "Artificial Intelligence and Games, Springer 2018  ·  Wurman et al., Nature 602:223–228, 2022",
        sz=10.5)
notes(s, """Trước khi nói về đồ án, em xin phác lại cách AI trong game đang được làm hiện nay.

Trong phần lớn dự án game, hành vi của NPC được lập trình bằng tay. Người phát triển viết máy trạng
thái, cây hành vi, hoặc đơn giản là một chuỗi luật if – else: thấy người chơi trong bán kính này thì
đuổi, máu xuống dưới ngưỡng kia thì bỏ chạy, tới điểm tuần tra thì quay đầu. Học tăng cường gần như
chưa xuất hiện trong quy trình sản xuất thông thường. Em xin nói ngay là có ngoại lệ đáng kể:
GT Sophy của Sony AI, công bố trên tạp chí Nature năm 2022, được huấn luyện bằng học tăng cường sâu
và nay đã thành tính năng chính thức trong Gran Turismo 7. Nhưng đó là dự án của một phòng nghiên
cứu lớn, chưa phải quy trình phổ biến.

Cách làm đó có lý do chính đáng. Logic viết tay chạy nhẹ, kết quả lặp lại được nên kiểm thử được, và
khi có lỗi thì người phát triển biết chính xác dòng nào gây ra. Với phần lớn NPC trong game thương
mại, chừng đó là đủ.

Nhưng nó có một cái giá. Hành vi của NPC chỉ phong phú đúng bằng số tình huống mà người viết đã
lường trước, nên nhân vật thường hành động cứng nhắc: lặp đi lặp lại một kịch bản, phản ứng giống
hệt nhau trong những tình huống giống nhau. Người chơi chỉ cần vài lượt là đoán được quy luật, rồi
khai thác chính quy luật đó. Muốn phong phú hơn thì phải viết thêm nhánh, và bộ luật phình ra rất
nhanh cho tới lúc không ai còn kiểm soát nổi.

Học tăng cường đi theo hướng ngược lại. Người phát triển không liệt kê tình huống nữa, mà định nghĩa
mục tiêu bằng một tín hiệu phần thưởng; hành vi nảy sinh từ quá trình tác nhân tự tương tác với môi
trường. Sơ đồ bên phải là vòng lặp đó: tác nhân chọn hành động, môi trường trả về trạng thái mới và
phần thưởng, rồi lặp lại hàng triệu lần.

Điều em muốn hội đồng lưu ý là cái khó không biến mất, nó chuyển chỗ. Thay vì viết luật, người phát
triển phải đặc tả đúng ba ô trong sơ đồ: tác nhân quan sát được gì, làm được gì, và được thưởng bao
nhiêu. Riêng ô phần thưởng là tín hiệu duy nhất định nghĩa mục tiêu — thuật toán không có cách nào
biết ta thật sự muốn gì ngoài con số đó. Và một đặc tả sai thì không hề báo lỗi: chương trình vẫn
chạy, TensorBoard vẫn vẽ đường cong, tác nhân chỉ đơn giản là học một hành vi khác với hành vi ta
muốn.

Hai con số bên trái là tỉ lệ công sức đo trên chính đồ án này: phần đặc tả ba môi trường chiếm hơn
ba mươi trang thuyết minh, còn cấu hình cho cả năm thuật toán gói gọn trong ba trang YAML. Đó cũng
là lý do đồ án của em đặt trọng tâm vào khâu xây dựng và kiểm chứng môi trường.

[Nguồn ghi ở chân slide, phòng khi hội đồng hỏi: cây hành vi do Damian Isla trình bày tại GDC 2005
cho hệ thống AI của Halo 2; giáo trình Artificial Intelligence and Games của Yannakakis và Togelius,
Springer 2018, xếp máy trạng thái và cây hành vi vào nhóm "soạn hành vi thủ công"; GT Sophy công bố
trên Nature tập 602, trang 223–228, năm 2022.]""")

# ==================================================== 3. MỤC TIÊU & ĐÓNG GÓP ===
s = new_slide_t("Mục tiêu và đóng góp")

_goals = [("1", "Xây dựng ba môi trường Unity ML-Agents", "đối kháng · hợp tác · điều hướng"),
          ("2", "Kiểm chứng bằng năm thuật toán", "PPO · SAC · MA-POCA · MAPPO · DQN"),
          ("3", "Đóng gói để dùng lại được", "gói UPM kèm 15 cấu hình đã chạy thật"),
          ("4", "Xây ứng dụng demo trực quan", "nạp 15 mô hình ONNX, ba chế độ chơi")]
for i, (n, h, b) in enumerate(_goals):
    x = 0.55 + (i % 2) * 6.33
    y = 1.30 + (i // 2) * 2.06
    card(s, x, y, 5.92, 1.86)
    oval(s, x + 0.28, y + 0.71, 0.44, fill=ORANGE, label=n, sz=15)
    textbox(s, x + 0.90, y + 0.16, 4.74, 1.54, [
        {"runs": [(h, {"sz": 16.5, "b": True, "c": NAVY})], "line": 1.04, "space_after": 5},
        {"runs": [(b, {"sz": 14, "c": MUTED})], "line": 1.06}], anchor="c")

card(s, 0.55, 5.52, 12.25, 1.32)
textbox(s, 0.86, 5.52, 11.63, 1.32, [
    {"runs": [("Phạm vi:  ", {"sz": 15, "b": True, "c": ORANGE}),
              ("ba môi trường Unity 3D; năm thuật toán đóng vai phép thử kiểm chứng môi trường.",
               {"sz": 15})], "line": 1.10, "space_after": 6},
    {"runs": [("Không bao gồm:  ", {"sz": 15, "b": True, "c": ORANGE}),
              ("đồ họa, âm thanh, sản phẩm game thương mại; đồ án không đề xuất thuật toán mới.",
               {"sz": 15})], "line": 1.10}], anchor="c")
notes(s, """Đồ án đặt ra bốn mục tiêu, và điểm em muốn nhấn mạnh là mỗi mục tiêu đều gắn với một
sản phẩm bàn giao cụ thể chứ không dừng ở mô tả.

Mục tiêu một là xây dựng ba môi trường Unity ML-Agents thuộc ba dạng bài toán khác nhau: đối
kháng, hợp tác và điều hướng đơn tác nhân. Em chọn ba dạng lệch hẳn nhau, chứ không phải ba biến
thể của cùng một dạng, vì mục đích là dò xem thiết kế môi trường ảnh hưởng thế nào tới kết quả.

Mục tiêu hai là kiểm chứng ba môi trường đó bằng năm thuật toán khác nguyên lý. Xin phép nói rõ
ngay ở đây, vì đây là điểm dễ hiểu nhầm nhất của đồ án: năm thuật toán không phải đối tượng đánh
giá. Chúng đóng vai năm phép thử độc lập. Nếu một môi trường được đặc tả tốt thì nó phải phản ứng
nhất quán trước nhiều hướng tiếp cận khác nhau; còn nếu cả năm cùng thất bại thì lỗi nằm ở đặc tả.

Mục tiêu ba là đóng gói để người khác dùng lại được, dưới dạng một gói Unity Package Manager kèm
mười lăm tệp cấu hình đã chạy thật. Mục tiêu bốn là ứng dụng demo nạp trực tiếp mười lăm mô hình
ONNX đã huấn luyện.

Về phạm vi, đồ án không làm đồ họa, không làm âm thanh, không hướng tới một sản phẩm game thương
mại, và cũng không đề xuất thuật toán học tăng cường mới.""")

# ======================================================= 4. KIẾN TRÚC TỔNG THỂ =
s = new_slide_t("Kiến trúc tổng thể — một đường ống đo duy nhất")

_stages = [("MÔI TRƯỜNG UNITY", "quan sát · hành động · reward"),
           ("ML-AGENTS TRAINER", "năm thuật toán phép thử"),
           ("TENSORBOARD", "reward · entropy · độ dài"),
           ("XUẤT ONNX", "15 mô hình đã huấn luyện"),
           ("ỨNG DỤNG DEMO", "ba chế độ · hai tab số liệu")]
_sw, _sgap = 2.21, 0.29
for i, (h, b) in enumerate(_stages):
    x = 0.55 + i * (_sw + _sgap)
    fill = NAVY if i in (0, 4) else CARD
    round_rect(s, x, 1.40, _sw, 1.70, fill=fill, adj=0.09)
    textbox(s, x + 0.10, 1.52, _sw - 0.20, 1.46, [
        {"runs": [(h, {"sz": 13, "b": True, "c": WHITE if i in (0, 4) else NAVY})],
         "align": "c", "line": 1.02, "space_after": 6},
        {"runs": [(b, {"sz": 11, "c": ICE if i in (0, 4) else MUTED})],
         "align": "c", "line": 1.02}], anchor="c")
    if i < 4:
        arrow(s, x + _sw + 0.03, 2.11, 0.23, 0.28, MSO_SHAPE.RIGHT_ARROW)

table(s, 0.55, 3.42, 12.25, 1.98, [
    ["Thuật toán", "Nguồn trainer", "Ghi chú triển khai"],
    ["PPO · SAC · MA-POCA", "có sẵn trong ML-Agents 4.0.2", "chạy trực tiếp bằng tệp cấu hình"],
    ["MAPPO", "trainer mở rộng do đồ án bổ sung", "phê bình tập trung, đăng ký thêm theo tên"],
    ["DQN", "trainer mở rộng do đồ án bổ sung", "chỉ nhận hành động rời rạc"],
], col_w=[2.5, 3.1, 4.0], hsz=13, bsz=13.5, row_h=0.46, align_first="l")

card(s, 0.55, 5.62, 12.25, 1.24)
textbox(s, 0.86, 5.62, 11.63, 1.24, [
    {"runs": [("Cả năm thuật toán dùng chung một đường ống: cùng gRPC, cùng TensorBoard, "
               "cùng đường xuất ONNX.", {"sz": 15, "b": True, "c": NAVY})],
     "line": 1.10, "space_after": 6},
    {"runs": [("Nhờ vậy chênh lệch đo được mới nói về môi trường, không phải về cách đo.",
               {"sz": 14, "c": INK})], "line": 1.08}], anchor="c")
notes(s, """Kiến trúc hệ thống đi theo một đường thẳng. Môi trường Unity sinh quan sát và nhận
hành động. Trainer của ML-Agents huấn luyện qua giao thức gRPC. TensorBoard ghi lại mọi chỉ số.
Mô hình được xuất ra định dạng ONNX. Và ứng dụng demo nạp thẳng tệp ONNX đó.

Bảng ở giữa là điểm cần nói thêm về mặt kỹ thuật. Gói Unity ML-Agents bản 4.0.2, cùng với phía
Python đi kèm, chỉ đăng ký sẵn ba trainer là ppo, poca và sac. Ba thuật toán này chạy trực tiếp bằng tệp cấu hình. Hai thuật toán
còn lại là MAPPO và DQN không có sẵn, nên em đã bổ sung chúng dưới dạng trainer mở rộng, đăng ký
thêm theo tên vào hệ thống trainer của thư viện.

Cách làm này quan trọng ở chỗ: hai trainer bổ sung dùng chung toàn bộ hạ tầng đo của ML-Agents —
cùng giao thức gRPC để nói chuyện với Unity, cùng cơ chế ghi TensorBoard, cùng đường xuất ONNX.
Nếu em viết riêng một vòng huấn luyện cho DQN thì mọi so sánh sau này sẽ vô nghĩa, vì không phân
biệt được chênh lệch đến từ thuật toán hay đến từ cách đo. Có chung một đường ống đo thì chênh
lệch quan sát được mới quy về môi trường được.

Một chi tiết nhỏ về DQN: thuật toán này chỉ làm việc với hành động rời rạc, nên riêng ở Football
Table em phải dựng thêm một phiên bản môi trường rời rạc hóa — em sẽ nói kỹ hơn ở phần thực nghiệm.""")

# ================================================== 5. BA MÔI TRƯỜNG ĐẠI DIỆN ==
s = new_slide_t("Ba môi trường đại diện ba dạng bài toán")

_envs = [("env_football.png", "Football Table", "ĐỐI KHÁNG 1-1"),
         ("env_puzzle.png", "Capture The Flag", "HỢP TÁC HAI TÁC NHÂN"),
         ("env_crossroad.png", "Cross The Road", "ĐIỀU HƯỚNG ĐƠN TÁC NHÂN")]
for i, (fn, name, kind) in enumerate(_envs):
    x = 0.55 + i * 4.15
    picture_plate(s, x, 1.22, 3.95, 2.22, os.path.join(IMG, fn), pad=0.06)
    simple(s, x, 3.50, 3.95, 0.30, name, sz=16, b=True, c=NAVY, align="c", wrap=False)
    simple(s, x, 3.82, 3.95, 0.26, kind, sz=11.5, b=True, c=ORANGE, align="c", wrap=False)

table(s, 0.55, 4.22, 12.25, 2.66, [
    ["", "Football Table", "Capture The Flag", "Cross The Road"],
    ["Quan sát", "60 giá trị", "61 giá trị", "132 giá trị"],
    ["Hành động", "liên tục, 8 chiều", "rời rạc, 7 hành động", "rời rạc, 4 hành động"],
    ["Cảm biến tia", "không", "7 tia · 6 nhãn", "21 tia · 4 nhãn"],
    ["Nhịp quyết định", "2 bước vật lý", "20 bước vật lý", "5 bước vật lý"],
    ["Phần thưởng nhóm", "không", "có", "không"],
    ["Cơ chế đặc thù", "tự chơi (self-play)", "bàn giao · học theo chương trình", "vật cản động"],
], col_w=[2.5, 3.1, 3.5, 3.1], hsz=13, bsz=13, row_h=0.38, align_first="l")
notes(s, """Ba môi trường được chọn để phủ ba dạng bài toán khác hẳn nhau.

Football Table là bàn bi lắc ba chiều, đối kháng một chọi một, hành động liên tục. Capture The
Flag là câu đố hợp tác: hai tác nhân phải bàn giao nhau mới mở được cửa. Cross The Road là điều
hướng đơn tác nhân qua năm làn xe, phần thưởng rất thưa.

Bảng phía dưới cho thấy ba môi trường lệch nhau trên nhiều trục cùng một lúc: cỡ không gian quan
sát chênh nhau hơn hai lần, dạng hành động khác nhau, có hay không dùng phần thưởng nhóm, và nhịp
ra quyết định chênh nhau tới mười lần.

Về nhịp ra quyết định, em xin giải thích thêm vì nó hay bị hỏi. Đây là tham số Decision Period của
thành phần DecisionRequester trong Unity: tác nhân chỉ hỏi mạng nơ-ron sau mỗi vài bước vật lý, và
giữ nguyên hành động cũ trong các bước còn lại. Một bước vật lý của Unity là hai phần trăm giây,
nên hai bước tương đương hai mươi lăm quyết định mỗi giây, còn hai mươi bước chỉ còn hai phẩy năm
quyết định mỗi giây. Nguyên tắc chọn là mỗi quyết định phải tương ứng với một hành vi trọn vẹn của
môi trường. Football Table điều khiển vật lý liên tục, bóng đi rất nhanh, nên phải bám gần khung
hình. Cross The Road đi theo ô, nhịp năm bước vừa đủ để cú trượt sang ô mới kết thúc trước khi hỏi
quyết định kế tiếp. Còn Capture The Flag là câu đố chân trời dài, một quyết định nên tương đương
"di chuyển tới một vị trí khác" chứ không phải "nhích một cái", nên nhịp thưa nhất.

Sự lệch đồng thời trên nhiều trục như vậy là chủ ý. Một bộ môi trường chỉ khác nhau ở đúng một
trục sẽ không đủ sức cho biết thuật toán nào thắng nhờ đâu.""")

# ================================================ 6. TRỌNG TÂM: HÀM PHẦN THƯỞNG
s = new_slide_t("Trọng tâm thiết kế — hàm phần thưởng")

card(s, 0.55, 1.24, 12.25, 1.94)
_bn = {"sz": 20, "b": True, "c": NAVY}
_sb = {"sz": 20, "b": True, "c": NAVY, "sub": True}
textbox(s, 0.80, 1.42, 11.75, 0.46, [
    {"runs": [("R(s, a, s′)  =  R", _bn), ("mt", _sb), ("(s′)  +  F(s, s′)  +  R", _bn),
              ("b", _sb), ("(s)", _bn)], "align": "c", "line": 1.0}], wrap=False)
_parts = [("Rmt · mục tiêu", "thưa, chỉ khác không tại cột mốc"),
          ("F · định hình thế năng", "F = γΦ(s′) − Φ(s) — bất biến chính sách (Ng và cs., 1999)"),
          ("Rb · theo bước", "dày — nơi phát sinh lợi dụng phần thưởng")]
for i, (h, b) in enumerate(_parts):
    x = 0.80 + i * 3.95
    textbox(s, x, 2.10, 3.75, 0.90, [
        {"runs": [(h, {"sz": 14, "b": True, "c": ORANGE})], "line": 1.02, "space_after": 4},
        {"runs": [(b, {"sz": 12.5, "c": INK})], "line": 1.06}], anchor="t")

_pr = [("1", "Bám theo lời giải mẫu", "đặt phần thưởng lên từng mốc của trình tự thao tác"),
       ("2", "Ưu tiên sự kiện và thế năng", "tránh cộng dồn theo bước — nơi sinh lối tắt"),
       ("3", "Giữ trật tự độ lớn", "tổng khoản phụ < phần thưởng hoàn thành nhiệm vụ"),
       ("4", "Sửa cơ chế trước khi tăng hệ số", "hệ số chỉ chỉnh sau khi cơ chế đã đúng")]
for i, (n, h, b) in enumerate(_pr):
    x = 0.55 + i * 3.11
    card(s, x, 3.44, 2.92, 2.30)
    oval(s, x + 0.22, 3.62, 0.36, fill=NAVY, label=n, sz=12.5)
    simple(s, x + 0.66, 3.58, 2.04, 0.46, h, sz=14, b=True, c=NAVY, line=1.02, anchor="c")
    simple(s, x + 0.22, 4.18, 2.48, 1.42, b, sz=12.5, c=INK, line=1.10, anchor="c")

card(s, 0.55, 5.98, 12.25, 0.90)
simple(s, 0.86, 5.98, 11.63, 0.90,
       "Đặc tả đi một chiều: hành vi mục tiêu → hàm phần thưởng → chính sách học được. "
       "Sai lệch chỉ lộ ra ở bước cuối.",
       sz=14.5, b=True, c=NAVY, align="c", anchor="c", wrap=False)
notes(s, """Hàm phần thưởng là thành phần khó thiết kế nhất, nên em xin dành riêng một slide.

Em tách hàm phần thưởng thành ba thành phần. Thành phần mục tiêu R chỉ số mt là phần thưa, chỉ
khác không tại các cột mốc — đây chính là thành phần định nghĩa bài toán. Thành phần F là định
hình theo thế năng, có dạng gamma nhân Phi của trạng thái mới trừ Phi của trạng thái cũ; dạng này
có một tính chất đã được chứng minh trong lý thuyết là không làm đổi lời giải tối ưu, nên dùng nó
để dẫn đường là an toàn. Thành phần R chỉ số b là phần theo bước, dày đặc, và đây chính là nơi
phát sinh hiện tượng lợi dụng phần thưởng.

Từ ba thành phần đó em rút ra bốn nguyên tắc. Nguyên tắc một: viết ra trước trình tự thao tác mà
một người chơi giỏi sẽ làm, rồi mới đặt phần thưởng lên từng mốc của trình tự đó — chứ không nghĩ
ra hệ số trước. Nguyên tắc hai: ưu tiên thưởng theo sự kiện hoặc theo thế năng, tránh cộng dồn
theo bước. Nguyên tắc ba: giữ trật tự độ lớn, tức tổng mọi khoản phụ trên trọn một lượt chơi phải
nhỏ hơn phần thưởng hoàn thành nhiệm vụ; nếu vi phạm thì tác nhân sẽ đi gom khoản phụ thay vì giải
nhiệm vụ. Nguyên tắc bốn: khi hành vi mong muốn không xuất hiện, hãy sửa cơ chế game hoặc chương
trình học trước, đừng vội nhân hệ số lên.

Dòng cuối slide là lý do vì sao việc này khó. Quy trình đặc tả đi một chiều: từ hành vi mục tiêu
trong đầu người thiết kế, sang hàm phần thưởng, rồi mới ra chính sách mà tác nhân thật sự học
được. Sai lệch chỉ lộ ra ở bước cuối cùng, và cách duy nhất để phát hiện là quan sát hành vi thật
rồi lần ngược về đặc tả. Ba slide tiếp theo cho thấy mỗi hằng số trong ba môi trường của em đều có
một lập luận định lượng đứng sau.""")

# ================================================== 7. FOOTBALL TABLE (THIẾT KẾ)
s = new_slide_t("Football Table — đối kháng 1-1, hành động liên tục")
picture_plate(s, 0.55, 1.24, 6.10, 3.62, os.path.join(IMG, "env_football.png"), pad=0.14)
caption(s, 0.55, 4.94, 6.10,
        "Bàn bi lắc 3D: bốn thanh gạt mỗi đội, mỗi thanh có hai bậc tự do là trượt và xoay.")

card(s, 0.55, 5.40, 6.10, 1.48)
textbox(s, 0.82, 5.40, 5.56, 1.48, [
    {"runs": [("Thách thức thiết kế", {"sz": 14.5, "b": True, "c": ORANGE})],
     "line": 1.0, "space_after": 5},
    {"runs": [("Gán công lao trong tự chơi: đối thủ mạnh dần theo chính mình, nên phần thưởng "
               "tuyệt đối không đo được tiến bộ.", {"sz": 13.5, "c": INK})], "line": 1.08}],
    anchor="c")

table(s, 6.95, 1.24, 5.85, 5.64, [
    ["Thành phần", "Đặc tả"],
    ["Quan sát", "60 giá trị · đối xứng hóa hai phía"],
    ["Hành động", "liên tục, 8 chiều"],
    ["Nhịp quyết định", "2 bước vật lý (25 qđ/s)"],
    ["Phần thưởng mục tiêu", "+5 ghi bàn  ·  −1 thủng lưới"],
    ["Phần thưởng theo bước", "−1/3000 mỗi bước"],
    ["Cơ chế", "tự chơi, đo bằng chỉ số ELO"],
    ["Kết thúc lượt", "có bàn thắng hoặc 3.000 bước"],
    ["Ngân sách", "1.600.000 bước"],
], col_w=[2.4, 3.5], hsz=13, bsz=13, row_h=0.60, align_first="l")
notes(s, """Football Table là bàn bi lắc ba chiều. Mỗi đội có bốn thanh gạt, mỗi thanh có hai bậc
tự do là trượt ngang và xoay, nên một tác nhân điều khiển tám chiều liên tục cùng lúc.

Quan sát gồm sáu mươi giá trị: trạng thái quả bóng, bốn thanh gạt đội nhà và bốn thanh gạt đối
thủ. Em đối xứng hóa quan sát để một mạng chính sách dùng chung được cho cả hai phía, nhờ vậy cơ
chế tự chơi mới hoạt động.

Bây giờ là phần em muốn giải thích kỹ, vì trên slide chỉ ghi con số. Vì sao thưởng ghi bàn
là cộng năm còn phạt thủng lưới chỉ trừ một, thay vì để hai giá trị bằng nhau? Hãy xét tình huống
tác nhân phải chọn giữa cầm cự — giữ bóng cho tới hết lượt chơi — và dứt điểm, tức tổ chức một pha
tấn công thành công với xác suất p và bị phản công thủng lưới với xác suất một trừ p. Cầm cự tới
hết ba nghìn bước thì gánh trọn khoản phạt thời gian, nên giá trị của nó bằng âm một. Dứt điểm có
giá trị kỳ vọng bằng p nhân phần thưởng ghi bàn, trừ đi một trừ p nhân hình phạt thủng lưới. So
sánh hai vế, tấn công chỉ đáng hơn cầm cự khi p vượt một ngưỡng bằng hình phạt thủng lưới trừ một,
chia cho tổng của thưởng và phạt. Với cấu hình đối xứng cộng năm trừ năm, ngưỡng ấy bằng không phẩy
bốn: tác nhân chỉ chịu tấn công khi tự tin thắng ít nhất bốn mươi phần trăm số pha bóng — điều gần
như không bao giờ đúng ở đầu huấn luyện — nên chính sách hội tụ về phòng ngự tuyệt đối và tín hiệu
học biến mất. Đổi sang cộng năm trừ một, ngưỡng tụt xuống đúng bằng không, nghĩa là mọi cơ hội ghi
bàn dù nhỏ đến mấy cũng đáng thử. Phép tính này còn cho thấy khoản phạt thời gian chẳng phải chi
tiết phụ, mà chính là thứ quyết định dấu của ngưỡng.

Con số một phần ba nghìn cũng không tùy tiện. Một lượt chơi tối đa ba nghìn bước, nên tổng phạt
thời gian trên trọn một lượt bằng đúng một điểm — bằng đúng hình phạt thủng lưới. Nói cách khác,
em định giá một trận hòa không bàn thắng ngang với thua một bàn, để tác nhân không coi việc kéo
dài thời gian là một chiến lược.

Thách thức lớn nhất ở môi trường này là gán công lao trong tự chơi: đối thủ mạnh dần lên theo
chính mình, nên phần thưởng tuyệt đối không phản ánh tiến bộ. Đó là lý do em phải dùng chỉ số ELO,
và em sẽ quay lại điểm này ở phần kết quả.""")

# ================================================ 8. CAPTURE THE FLAG (THIẾT KẾ)
s = new_slide_t("Capture The Flag — hợp tác hai tác nhân")
picture_plate(s, 0.55, 1.24, 6.10, 3.62, os.path.join(IMG, "env_puzzle.png"), pad=0.14)
caption(s, 0.55, 4.94, 6.10,
        "Hành lang ba khoang: hai bàn đạp nằm hai phía đối diện nhưng cùng điều khiển một cổng cửa.")

card(s, 0.55, 5.40, 6.10, 1.48)
textbox(s, 0.82, 5.40, 5.56, 1.48, [
    {"runs": [("Thách thức thiết kế", {"sz": 14.5, "b": True, "c": ORANGE})],
     "line": 1.0, "space_after": 5},
    {"runs": [("Gán công lao cho cả nhóm: phần thưởng chung phải khiến hợp tác có lợi mà không "
               "thưởng cho việc ăn theo.", {"sz": 13.5, "c": INK})], "line": 1.08}], anchor="c")

table(s, 6.95, 1.24, 5.85, 5.64, [
    ["Thành phần", "Đặc tả"],
    ["Quan sát", "61 giá trị · 5 cờ logic + 7 tia"],
    ["Hành động", "rời rạc, 7 hành động"],
    ["Nhịp quyết định", "20 bước vật lý (2,5 qđ/s)"],
    ["Phần thưởng nhóm", "+1 cả hai tới đích  ·  +0,5 vượt cổng"],
    ["Định hình thế năng", "κ = 0,01"],
    ["Khoản theo bước", "chuẩn hóa theo N = 100.000"],
    ["Kết thúc lượt", "cả hai tới đích hoặc 100.000 bước"],
    ["Ngân sách", "1.700.000 bước"],
], col_w=[2.4, 3.5], hsz=13, bsz=13, row_h=0.60, align_first="l")
notes(s, """Capture The Flag là câu đố bàn giao. Bản đồ là một hành lang ba khoang. Có hai bàn đạp
cùng điều khiển một cổng cửa, nhưng chúng nằm ở hai phía đối diện của cổng đó. Hệ quả là không ai
tự mình giải được: một tác nhân phải đứng lên bàn đạp giữ cửa mở cho tác nhân kia đi qua, rồi tới
lượt người vừa qua cửa giữ cửa cho người còn lại.

Quan sát sáu mươi mốt giá trị, trong đó năm cờ logic là mấu chốt. Chúng cho tác nhân biết bàn đạp
nào đang bị đè và ai đã vượt cổng — nếu thiếu thông tin này thì tác nhân không có cách nào biết
khi nào được phép rời bàn đạp.

Có hai chi tiết thiết kế em muốn giải thích thêm.

Thứ nhất là khoảng ân hạn của cửa. Khi tác nhân rời bàn đạp, cửa không sập ngay mà giữ mở thêm
một giây. Em phải chứng minh con số một giây này vừa đủ mà không quá tay. Khoảng cách từ bàn đạp
tới cổng là mười chín đơn vị. Vận tốc tối đa của tác nhân là mười bảy phẩy năm đơn vị mỗi
giây, suy ra từ thế cân bằng giữa lượng vận tốc cộng thêm mỗi bước vật lý là một phẩy bốn và hệ số
cản tuyến tính bằng bốn — lấy một phẩy bốn chia cho tích của bốn với hai phần trăm giây.
Nghĩa là ngay cả khi giả định tác nhân đã chạy sẵn ở vận tốc tối đa theo đúng hướng, trong một
giây nó vẫn chưa tới được cổng; còn nếu xuất phát từ trạng thái đứng yên thì chỉ đi được khoảng
mười ba phẩy hai đơn vị. Kết luận: không tồn tại quỹ đạo đơn độc nào vượt được cổng. Nhờ vậy mọi
kết quả cao trên môi trường này đều chắc chắn là hợp tác thật, chứ không phải một tác nhân tự lách
luật.

Thứ hai là việc chuẩn hóa. Mọi khoản thưởng theo bước — đứng bàn đạp, đẩy khối, phạt thời gian —
đều chia cho N bằng một trăm nghìn, đúng bằng số bước tối đa của một lượt chơi. Cách này bảo đảm
tổng mọi khoản phụ trên trọn một lượt luôn nhỏ hơn phần thưởng hoàn thành nhiệm vụ, đúng theo
nguyên tắc giữ trật tự độ lớn ở slide trước.

[Ghi chú dự phòng, chỉ dùng nếu hội đồng hỏi tới: bảng phần thưởng trong thuyết minh ghi thưởng
nhóm khi cả hai tới đích là +1,0, còn trong scene huấn luyện thực tế trường reward của trigger
checkpoint đang đặt bằng 2 ở cả tám khu vực. Đây là chênh lệch giữa tài liệu và cấu hình, không
ảnh hưởng tới số liệu đã báo cáo vì các con số ở Chương 4 lấy từ chỉ số phần thưởng cá nhân
(Environment/Cumulative Reward), còn khoản này là phần thưởng nhóm. Lập luận giữ trật tự độ lớn
cũng không đổi: giá trị càng lớn thì phần thưởng hoàn thành càng vượt xa tổng các khoản phụ. Em
xin ghi nhận và đính chính lại trong bản hoàn thiện.]""")

# ================================================== 9. CROSS THE ROAD (THIẾT KẾ)
s = new_slide_t("Cross The Road — điều hướng với phần thưởng thưa")
picture_plate(s, 0.55, 1.24, 6.10, 3.62, os.path.join(IMG, "env_crossroad.png"), pad=0.14)
caption(s, 0.55, 4.94, 6.10,
        "Năm làn xe; tốc độ mỗi xe bốc ngẫu nhiên lại sau mỗi lượt chơi nên không học vẹt được.")

card(s, 0.55, 5.40, 6.10, 1.48)
textbox(s, 0.82, 5.40, 5.56, 1.48, [
    {"runs": [("Thách thức thiết kế", {"sz": 14.5, "b": True, "c": ORANGE})],
     "line": 1.0, "space_after": 5},
    {"runs": [("Chống chính sách tự đóng băng: lượt chơi không giới hạn bước nên đứng yên mãi "
               "cũng là một lời giải hợp lệ.", {"sz": 13.5, "c": INK})], "line": 1.08}],
    anchor="c")

table(s, 6.95, 1.24, 5.85, 5.64, [
    ["Thành phần", "Đặc tả"],
    ["Quan sát", "132 giá trị · 21 tia cảm biến"],
    ["Hành động", "rời rạc, 4 hành động"],
    ["Nhịp quyết định", "5 bước vật lý (10 qđ/s)"],
    ["Phần thưởng mục tiêu", "+1 tới đích  ·  −0,025 va chạm"],
    ["Tín hiệu tò mò", "cường độ 0,02"],
    ["Kết thúc lượt", "tới đích hoặc va chạm"],
    ["Giới hạn bước", "không giới hạn"],
    ["Ngân sách", "2.000.000 bước"],
], col_w=[2.4, 3.5], hsz=13, bsz=13, row_h=0.60, align_first="l")
notes(s, """Cross The Road đơn giản về luật chơi nhưng khó về tín hiệu học. Tác nhân phải băng qua
năm làn xe để tới đích. Tốc độ mỗi xe được bốc ngẫu nhiên lại sau mỗi lượt chơi, nên tác nhân
không thể học vẹt một trình tự cố định mà buộc phải đọc tình huống.

Quan sát một trăm ba mươi hai giá trị: tọa độ của tác nhân và của đích cho biết đi về đâu, còn hai
mươi mốt tia cảm biến cho biết lúc nào an toàn. Bốn hành động rời rạc là đứng yên, sang trái, sang
phải và tiến lên. Xin lưu ý rằng đứng yên ở đây là một nước đi chiến thuật thật sự — chờ xe đi qua
— chứ không phải một lựa chọn thừa.

Điểm em muốn giải thích kỹ là mức phạt va chạm chỉ có không phẩy không hai lăm, nghe rất nhẹ so
với phần thưởng tới đích là một. Lý do nằm ở chỗ môi trường này không giới hạn số bước. Khi lượt
chơi không bị cắt theo thời gian, việc đứng yên vĩnh viễn luôn có giá trị kỳ vọng bằng đúng không
— tác nhân không mất gì cả. Bây giờ xét quyết định đi tiếp: gọi p là xác suất qua được an toàn.
Nếu phạt va chạm bằng một, tức cân bằng với thưởng, thì đi tiếp chỉ đáng khi p vượt năm mươi phần
trăm. Ở giai đoạn đầu huấn luyện điều đó không xảy ra, nên chính sách tự đóng băng: tác nhân đứng
im mãi mãi và không bao giờ thu được kinh nghiệm để cải thiện. Với mức phạt không phẩy không hai
lăm, ngưỡng đáng đi tiếp chỉ còn hai phẩy bốn phần trăm.

Điều đáng nói là răn đe thật ở môi trường này không nằm ở con số phạt, mà nằm ở chỗ va chạm kết
thúc lượt chơi ngay lập tức — tức mất trắng cơ hội nhận một điểm.

Em còn bổ sung tín hiệu tò mò với cường độ không phẩy không hai, bằng một phần năm mươi trọng số
tín hiệu ngoại vi, để tạo động lực khám phá mà không phải dùng tới phạt thời gian.""")

# ====================================================== 10. THIẾT KẾ THỰC NGHIỆM
s = new_slide_t("Thiết kế thực nghiệm — lưới đầy đủ 5 × 3")

for i, (num, lab) in enumerate([("15", "lần huấn luyện đầy đủ"), ("5", "thuật toán làm phép thử"),
                                ("3", "môi trường game 3D"), ("×8", "khu vực chạy song song")]):
    stat_tile(s, 0.55 + i * 3.11, 1.24, 2.92, 1.20, num, lab, nsz=30, lsz=12.5)

table(s, 0.55, 2.72, 7.15, 1.80, [
    ["Môi trường", "PPO", "SAC", "MA-POCA", "MAPPO", "DQN"],
    ["Football Table", "✓", "✓", "✓", "✓", "✓ *"],
    ["Capture The Flag", "✓", "✓", "✓", "✓", "✓"],
    ["Cross The Road", "✓", "✓", "✓", "✓", "✓"],
], col_w=[2.3, 1, 1, 1.3, 1.1, 1.1], hsz=12.5, bsz=13, row_h=0.42, align_first="l")
caption(s, 0.55, 4.60, 7.15,
        "* DQN chỉ nhận hành động rời rạc nên Football Table cần một bản dựng riêng, chạy không tự chơi.")

table(s, 7.95, 2.72, 4.85, 2.30, [
    ["Ngân sách bước", ""],
    ["Cross The Road", "2.000.000"],
    ["Capture The Flag", "1.700.000"],
    ["Football Table", "1.600.000"],
    ["Phần cứng", "i5-14600KF · RTX 3050"],
], col_w=[2.6, 2.25], hsz=12.5, bsz=13, row_h=0.42, align_first="l")

card(s, 0.55, 5.22, 12.25, 1.66)
textbox(s, 0.86, 5.22, 11.63, 1.66, [
    {"runs": [("Trong từng môi trường, cả năm thuật toán dùng chung một ngân sách bước và cùng "
               "phần cứng.", {"sz": 15, "b": True, "c": NAVY})], "line": 1.10, "space_after": 7},
    {"runs": [("Một môi trường chỉ được coi là kiểm chứng chắc chắn khi nó phản ứng nhất quán "
               "trước nhiều hướng tiếp cận khác nguyên lý.", {"sz": 14.5, "c": INK})],
     "line": 1.08}], anchor="c")
notes(s, """Thiết kế thực nghiệm là một lưới đầy đủ: năm thuật toán nhân ba môi trường, tổng cộng
mười lăm lần huấn luyện, không bỏ trống ô nào.

Trong từng môi trường, cả năm thuật toán dùng chung một ngân sách bước và chạy trên cùng một cấu
hình phần cứng. Ngân sách khác nhau giữa các môi trường vì độ khó khác nhau: Cross The Road hai
triệu bước, Capture The Flag một triệu bảy, Football Table một triệu sáu. Mỗi lần chạy dùng tám
khu vực huấn luyện song song trong cùng một cảnh Unity, nhờ vậy thu được nhiều quỹ đạo hơn trong
cùng thời gian thực.

Dấu sao ở ô DQN của Football Table cần giải thích. DQN chỉ làm việc với hành động rời rạc, trong
khi Football Table vốn là bài toán liên tục tám chiều. Em phải dựng một phiên bản rời rạc hóa
riêng cho lần chạy này, và phiên bản đó không chạy được cơ chế tự chơi. Hệ quả là DQN trên
Football Table được huấn luyện trong điều kiện khác hẳn bốn thuật toán còn lại và không có chỉ số
ELO. Em xin nói rõ để hội đồng lưu ý: số liệu của ô đó chỉ nên xem như một tham chiếu độc lập,
không đặt cạnh bốn thuật toán kia để xếp hạng.

Câu ở dưới cùng là tinh thần của toàn bộ thiết kế thực nghiệm này. Em không so sánh thuật toán để
tìm ra thuật toán tốt nhất. Em dùng năm thuật toán như năm phép thử, và chỉ khi một môi trường
phản ứng nhất quán trước cả năm hướng tiếp cận khác nguyên lý thì em mới dám nói môi trường đó đã
được kiểm chứng.""")

# ================================================= 11. TIÊU CHÍ NGHIỆM THU =====
s = new_slide_t("Tiêu chí nghiệm thu môi trường")

_crit = [("①", "HỌC ĐƯỢC", "≥ 4/5 thuật toán tiến bộ rõ rệt so với lúc khởi tạo"),
         ("②", "PHÂN BIỆT ĐƯỢC", "các hướng tiếp cận khác nguyên lý tách được nhau trên chỉ số đo"),
         ("③", "KHÔNG CÓ LỐI TẮT", "không lần chạy nào đạt điểm sát trần mà chưa giải được nhiệm vụ")]
for i, (num, h, b) in enumerate(_crit):
    x = 0.55 + i * 4.15
    card(s, x, 1.24, 3.95, 2.16)
    oval(s, x + 0.26, 1.44, 0.42, fill=ORANGE, label=num, sz=15)
    simple(s, x + 0.80, 1.48, 2.90, 0.36, h, sz=16, b=True, c=NAVY, wrap=False)
    simple(s, x + 0.26, 1.98, 3.43, 1.26, b, sz=13.5, c=INK, line=1.10, anchor="c")

table(s, 0.55, 3.62, 12.25, 1.72, [
    ["Chỉ số đo", "Định nghĩa"],
    ["Mức năng lực", "phần thưởng trung bình trên 10% cuối của quá trình huấn luyện"],
    ["Tốc độ hội tụ", "bước mà từ đó đường cong không tụt dưới 90% mức ổn định"],
    ["Độ ổn định", "độ lệch chuẩn của đường cong trên 10% cuối"],
], col_w=[3.0, 9.25], hsz=13, bsz=13, row_h=0.43, align_first="l")

card(s, 0.55, 5.54, 12.25, 1.34)
textbox(s, 0.86, 5.54, 11.63, 1.34, [
    {"runs": [("Kết quả nghiệm thu:  ", {"sz": 15, "b": True, "c": ORANGE}),
              ("Cross The Road và Capture The Flag đạt cả ba tiêu chí; Football Table đạt 2/3 "
               "do dải ELO quá hẹp.", {"sz": 15})], "line": 1.10, "space_after": 6},
    {"runs": [("Football Table dùng thêm chỉ số ELO vì trong tự chơi phần thưởng không so trực "
               "tiếp được.", {"sz": 14, "c": MUTED})], "line": 1.08}], anchor="c")
notes(s, """Để khẳng định một môi trường là dùng được, em đặt ra ba tiêu chí nghiệm thu và bốn chỉ
số đo, tất cả đều định lượng và đều xác định trước khi nhìn kết quả.

Tiêu chí một là học được: ít nhất bốn trên năm thuật toán phải tiến bộ rõ rệt so với lúc khởi tạo.
Nếu không thuật toán nào học được thì lỗi nằm ở đặc tả môi trường chứ không ở thuật toán.

Tiêu chí hai là phân biệt được: các hướng tiếp cận khác nguyên lý phải tách được nhau trên chỉ số
đo. Một môi trường cho mọi thuật toán cùng một điểm số thì không dùng để so sánh được, vì nó không
có độ phân giải.

Tiêu chí ba là không có lối tắt: không lần chạy nào được đạt điểm sát trần mà chưa thật sự giải
được nhiệm vụ. Đây là tiêu chí bắt lỗi lợi dụng phần thưởng, và cách kiểm tra là đối chiếu phần
thưởng với độ dài lượt chơi cùng hành vi quan sát được trong ứng dụng demo.

Bảng ở giữa là định nghĩa của các chỉ số đo. Mức năng lực lấy trung bình trên mười phần trăm cuối
chứ không lấy điểm cuối, để con số không phụ thuộc vào một bản ghi đơn lẻ. Ngoại lệ duy nhất là
ELO: đại lượng này tích lũy theo thời gian nên chỉ giá trị tại bước cuối mới có nghĩa.

Kết quả nghiệm thu: Cross The Road và Capture The Flag đạt cả ba tiêu chí. Football Table chỉ đạt
hai trên ba — nó học được và không có lối tắt, nhưng dải ELO thu được quá hẹp nên tiêu chí phân
biệt chưa đạt. Em sẽ trình bày rõ ở phần kết quả và cũng nêu lại trong phần hạn chế.""")

# ==================================================== 12. KẾT QUẢ TỔNG HỢP =====
s = new_slide_t("Kết quả — 15 lần huấn luyện")

table(s, 0.55, 1.26, 12.25, 3.06, [
    ["Thuật toán", "Cross The Road\nreward 10% cuối", "Capture The Flag\nreward 10% cuối",
     "Football Table\nELO cuối (mức tăng)"],
    ["PPO", "0,690", "0,834", "1259  (+60)"],
    ["SAC", "0,846", "0,183", "1254  (+57)"],
    ["MA-POCA", "0,556", "0,838", "1237  (+36)"],
    ["MAPPO", "0,835", "0,840", "1260  (+64)"],
    ["DQN", "0,336", "0,800", "— không tự chơi"],
], col_w=[2.4, 3.3, 3.3, 3.25], hsz=13, bsz=14, row_h=0.44)

_read = [("Cross The Road phân biệt tốt nhất", "phổ trải 0,336 – 0,846; 9/10 cặp tách được nhau"),
         ("Capture The Flag tách nhóm giải được", "bốn thuật toán 0,800 – 0,840, riêng SAC 0,183"),
         ("Football Table học được nhưng hẹp", "cả bốn đường ELO đều lên, dải chỉ 1237 – 1260")]
for i, (h, b) in enumerate(_read):
    x = 0.55 + i * 4.15
    headline(s, x, 4.34, 3.95, 1.40, h, b, hsz=14.5, bsz=12.5)

card(s, 0.55, 5.96, 12.25, 0.92)
simple(s, 0.86, 5.96, 11.63, 0.92,
       "Mọi số liệu trích thẳng từ tệp sự kiện TensorBoard của 15 lần chạy, không hiệu chỉnh.",
       sz=14, b=True, c=NAVY, align="c", anchor="c", wrap=False)
notes(s, """Đây là toàn bộ kết quả của mười lăm lần huấn luyện, gói trong một bảng.

Hai cột giữa là phần thưởng trung bình trên mười phần trăm cuối của Cross The Road và Capture The
Flag. Cột cuối là chỉ số ELO của Football Table, kèm mức tăng so với điểm xuất phát của chính lần
chạy đó.

Xin lưu ý một điểm về cách đọc cột ELO. Chỉ số ELO của hai lần chạy khác nhau không đem so trực
tiếp với nhau được, bởi mỗi lần chạy tính điểm dựa trên kho đối thủ của riêng nó. Đại lượng có
nghĩa ở đây là mức tăng so với giá trị khởi tạo, chứ không phải con số tuyệt đối. Ô DQN để trống
vì lần chạy đó không dùng tự chơi.

Ba nhận định rút ra từ bảng này. Cross The Road là môi trường phân biệt tốt nhất: phổ kết quả trải
từ không phẩy ba ba sáu tới không phẩy tám bốn sáu, và chín trên mười cặp thuật toán tách được
nhau. Capture The Flag tách rất rõ giữa nhóm giải được và nhóm thất bại: bốn thuật toán nằm gọn
trong dải không phẩy tám tới không phẩy tám bốn, riêng SAC dừng ở không phẩy một tám ba. Football
Table thì cả bốn đường ELO đều đi lên, nghĩa là môi trường học được, nhưng dải kết quả từ một
nghìn hai trăm ba bảy tới một nghìn hai trăm sáu mươi quá hẹp để xếp hạng.

Em xin nhấn mạnh dòng cuối: mọi con số trong bảng đều trích thẳng từ tệp sự kiện TensorBoard của
mười lăm lần chạy, không qua bất kỳ hiệu chỉnh nào. Ba slide tiếp theo là ảnh chụp trực tiếp từ
TensorBoard để hội đồng đối chiếu.""")

# ============================================= 13. FOOTBALL TABLE — ĐƯỜNG ELO ==
s = new_slide_t("Football Table — chỉ số ELO trong tự chơi")
tb_plate(s, 0.55, 1.24, 7.35, "tb_football_elo.png")
caption(s, 0.55, 6.62, 7.35,
        "Ảnh chụp trực tiếp từ TensorBoard · làm mượt 0,6 · DQN chạy không tự chơi nên không có ELO.")

table(s, 8.25, 1.24, 4.55, 1.98, [
    ["Thuật toán", "ELO đầu → cuối", "Tăng"],
    ["MAPPO", "1196 → 1260", "+64"],
    ["PPO", "1199 → 1259", "+60"],
    ["SAC", "1197 → 1254", "+57"],
    ["MA-POCA", "1201 → 1237", "+36"],
], col_w=[1.7, 1.9, 0.95], hsz=12, bsz=12.5, row_h=0.38, align_first="l")

note_line(s, 8.25, 3.46, 4.55, 1.62, 1, "Phần thưởng không xếp cùng thứ hạng với ELO",
          "Trong tự chơi, phần thưởng phụ thuộc vào việc gặp phiên bản đối thủ nào.")
note_line(s, 8.25, 5.26, 4.55, 1.62, 2, "Cả bốn đều lên, nhưng dải quá hẹp",
          "Khoảng 1237–1260 hẹp hơn dao động của chính đường cong nên chưa xếp hạng được.")
notes(s, """Đây là ảnh chụp trực tiếp từ TensorBoard, không phải biểu đồ em vẽ lại. Trục ngang là
số bước huấn luyện, trục dọc là chỉ số ELO, mức làm mượt để ở không phẩy sáu — đường đậm là đường
đã làm mượt, đường nhạt phía sau là số liệu thô.

Trước hết xin nói vì sao slide này không có biểu đồ phần thưởng. Trong cơ chế tự chơi, đối thủ
chính là các phiên bản trước của bản thân tác nhân. Khi tác nhân mạnh lên thì đối thủ cũng mạnh
lên, nên phần thưởng tuyệt đối dao động quanh một mức gần như không đổi thay vì tăng đều. Nói cách
khác, phần thưởng ở đây đo độ chênh lệch giữa tác nhân và bản sao của nó, chứ không đo năng lực.
Thứ hạng theo phần thưởng cũng không trùng với thứ hạng theo ELO, điều đó càng khẳng định phần
thưởng không phải thước đo đáng tin trong tự chơi.

ELO thì khác. ML-Agents tính riêng cho từng lần chạy, dựa trên kết quả đối đầu giữa chính sách
hiện tại và kho các phiên bản trước đó. Đại lượng có nghĩa là mức tăng so với điểm xuất phát.
Chiếu theo tiêu chí đó, cả bốn thuật toán đều tiến bộ: MAPPO tăng khoảng sáu mươi tư điểm, PPO sáu
mươi, SAC năm mươi bảy, MA-POCA ba mươi sáu. Không lần chạy nào kết thúc dưới điểm xuất phát, nên
kết luận là môi trường này học được.

Nhìn vào hình, diễn biến các đường lại rất khác nhau. Đường PPO tăng gần như đơn điệu tới khoảng
một triệu một trăm nghìn bước rồi nằm ngang hẳn — dấu hiệu cho thấy chính sách của nó không còn
vượt nổi kho đối thủ nữa. MAPPO thì lên chậm hơn nhưng vẫn còn xu hướng tăng ở cuối ngân sách.

Điều em phải nói thẳng là dải kết quả từ một nghìn hai trăm ba bảy tới một nghìn hai trăm sáu mươi
còn hẹp hơn biên độ dao động của chính các đường cong. Vì vậy em kết luận môi trường học được,
nhưng không xếp hạng nhóm dẫn đầu. Đây cũng là lý do Football Table chỉ đạt hai trên ba tiêu chí
nghiệm thu.""")

# ============================================ 14. CAPTURE THE FLAG — ĐƯỜNG CONG
s = new_slide_t("Capture The Flag — đường cong phần thưởng")
tb_plate(s, 0.55, 1.24, 7.35, "tb_ctf_reward.png")
caption(s, 0.55, 6.62, 7.35,
        "Ảnh chụp trực tiếp từ TensorBoard · làm mượt 0,6 · SAC hiện thành hai đoạn do có lần chạy tiếp.")

table(s, 8.25, 1.24, 4.55, 2.32, [
    ["Thuật toán", "Reward", "Hội tụ", "Độ dài lượt"],
    ["MAPPO", "0,840", "~260K", "153,3"],
    ["MA-POCA", "0,838", "~540K", "174,3"],
    ["PPO", "0,834", "~180K", "292,9"],
    ["DQN", "0,800", "~1,04M", "2017,9"],
    ["SAC", "0,183", "không đạt", "4999,0"],
], col_w=[1.5, 1.0, 1.1, 1.25], hsz=11.5, bsz=12.5, row_h=0.38, align_first="l")

note_line(s, 8.25, 3.80, 4.55, 1.44, 1, "Độ dài lượt là bằng chứng chắc hơn phần thưởng",
          "Bốn thuật toán dẫn đầu kéo lượt chơi xuống; SAC nằm ở trần 4999.")
note_line(s, 8.25, 5.44, 4.55, 1.44, 2, "Phê bình tập trung không tạo khác biệt đo được",
          "PPO 0,834 so với MAPPO 0,840 và MA-POCA 0,838 — chênh dưới 1%.")
notes(s, """Đường cong Capture The Flag tách rất rõ giữa nhóm giải được và nhóm thất bại. Bốn
đường leo lên dải không phẩy tám tới không phẩy tám bốn rồi nằm yên, cho thấy chúng đã tìm ra
trình tự bàn giao và giữ được nó.

Xin giải thích về đường SAC, vì trên hình nó hiện thành hai đoạn tách rời với chú giải SAC hai
chấm không và SAC hai chấm một. Lần chạy SAC này có một lần chạy tiếp từ điểm lưu tại khoảng một
triệu không tám trăm nghìn bước, nên trục bước không đơn điệu. Em bật đúng tùy chọn "tách trục
không đơn điệu" của TensorBoard để nó vẽ thành hai đoạn thay vì nối thẳng hai đầu — nếu không thì
sẽ có một đường chéo giả nối ngược. Đoạn đầu là lần chạy gốc, dao động rất mạnh và chỉ lên tới
khoảng không phẩy năm tám. Đoạn sau là lần chạy tiếp, sụp xuống quanh không phẩy không tám và nằm
đó cho tới hết ngân sách. Con số không phẩy một tám ba trong bảng là trung bình mười phần trăm
cuối của chuỗi đã gộp, tức phản ánh đoạn sụp này.

Điểm em muốn nhấn mạnh là cột độ dài lượt chơi, vì nó là bằng chứng chắc hơn cả phần thưởng. Bốn
thuật toán dẫn đầu đều kéo độ dài lượt chơi xuống rõ rệt — MAPPO còn một trăm năm ba bước,
MA-POCA một trăm bảy tư, PPO hai trăm chín ba. Muốn kết thúc lượt chơi nhanh thì bắt buộc phải
giải xong câu đố, nên đây là bằng chứng trực tiếp cho việc không có lối tắt. Ngược lại SAC nằm ở
trần bốn nghìn chín trăm chín chín gần như suốt quá trình, nghĩa là hầu hết các lượt đều hết giờ
mà chưa tới đích.

Nhận xét cuối: trên bài toán hợp tác có hai tác nhân đồng nhất, phê bình tập trung của MAPPO và
MA-POCA không tạo nổi khác biệt đo được so với PPO thường — chênh lệch nằm dưới một phần trăm, tức
nhỏ hơn dao động của chính đường cong.""")

# =============================================== 15. CROSS THE ROAD — ĐƯỜNG CONG
s = new_slide_t("Cross The Road — đường cong phần thưởng")
tb_plate(s, 0.55, 1.24, 7.35, "tb_ctr_reward.png")
caption(s, 0.55, 6.62, 7.35,
        "Ảnh chụp trực tiếp từ TensorBoard · làm mượt 0,6 · ngân sách 2.000.000 bước.")

table(s, 8.25, 1.24, 4.55, 2.32, [
    ["Thuật toán", "Reward", "Hội tụ", "Độ dài lượt"],
    ["SAC", "0,846", "~1,79M", "61,8"],
    ["MAPPO", "0,835", "~1,72M", "52,0"],
    ["PPO", "0,690", "~1,97M", "811,6"],
    ["MA-POCA", "0,556", "~450K", "42,2"],
    ["DQN", "0,336", "không đạt", "124,7"],
], col_w=[1.5, 1.0, 1.1, 1.25], hsz=11.5, bsz=12.5, row_h=0.38, align_first="l")

note_line(s, 8.25, 3.80, 4.55, 1.44, 1, "Hội tụ sớm không đồng nghĩa học nhanh",
          "MA-POCA ổn định từ ~450K nhưng chốt ở 0,556 rồi nằm im.")
note_line(s, 8.25, 5.44, 4.55, 1.44, 2, "MAPPO vượt PPO ở môi trường ĐƠN tác nhân",
          "0,835 so với 0,690 — đồ án chưa giải thích được cơ chế.")
notes(s, """Cross The Road là môi trường phân biệt tốt nhất trong ba môi trường của đồ án. Cả năm
đường đều tiến bộ rõ so với lúc khởi tạo, nên tiêu chí học được đạt. Phổ kết quả trải từ không
phẩy ba ba sáu của DQN lên tới không phẩy tám bốn sáu của SAC, và chín trên mười cặp thuật toán
tách được nhau — đó là lý do em nói môi trường này cho phép so sánh sạch nhất.

Nhìn vào hình có thể thấy hai nhóm hành vi rất khác nhau. Nhóm PPO, MA-POCA và MAPPO tăng ngay từ
đầu. Nhóm SAC và DQN gần như nằm im ở mức không suốt sáu trăm nghìn bước đầu, thậm chí có đoạn
xuống dưới không, rồi mới bật lên rất nhanh. Đây là biểu hiện điển hình của bài toán phần thưởng
thưa: tác nhân phải khám phá đủ lâu mới chạm được tín hiệu đầu tiên, và sau đó thì học rất nhanh.

Hai điểm em muốn nêu từ bảng bên phải.

Điểm một: hội tụ sớm không đồng nghĩa với học nhanh. MA-POCA ổn định sớm nhất, chỉ khoảng bốn trăm
năm mươi nghìn bước, nhưng chốt ở không phẩy năm năm sáu rồi nằm im tới hết. Đó là sự ổn định của
một chính sách tầm thường, không phải dấu hiệu tốt. Nếu chỉ đọc cột tốc độ hội tụ mà bỏ cột mức
năng lực thì sẽ kết luận ngược hẳn.

Điểm hai là một hiện tượng em chưa giải thích được và xin nêu trung thực. MAPPO đạt không phẩy tám
ba lăm, vượt hẳn PPO ở mức không phẩy sáu chín, ngay tại một môi trường chỉ có một tác nhân. Về
nguyên lý, phê bình tập trung của MAPPO lẽ ra vô hiệu khi nhóm chỉ có một thành viên, vì lúc đó nó
suy biến về đúng PPO. Em nghi ngờ khác biệt đến từ chi tiết triển khai trong trainer mở rộng hoặc
từ may rủi của một lần khởi tạo, nhưng chưa có bằng chứng, nên em ghi nó vào phần việc cần kiểm
chứng tiếp.""")

# ================================================== 16. LUẬN ĐIỂM RÚT RA =======
s = new_slide_t("Luận điểm rút ra từ kết quả")

card(s, 0.55, 1.30, 6.05, 2.72)
simple(s, 0.81, 1.50, 5.53, 0.32, "CÙNG MỘT THUẬT TOÁN, HAI MÔI TRƯỜNG",
       sz=12.5, b=True, c=ORANGE, wrap=False)
simple(s, 0.81, 1.88, 5.53, 0.30, "SAC trên Cross The Road và Capture The Flag",
       sz=13, c=MUTED, wrap=False)
textbox(s, 0.81, 2.26, 5.53, 0.80, [
    {"runs": [("0,846", {"sz": 40, "b": True, "c": NAVY}),
              ("     →     ", {"sz": 26, "c": MUTED}),
              ("0,183", {"sz": 40, "b": True, "c": RED})], "align": "c", "line": 1.0}])
simple(s, 0.81, 3.14, 5.53, 0.72,
       "Cùng thuật toán, cùng phần cứng, cùng đường ống đo. Chênh lệch 0,663 sinh ra từ đặc "
       "tính bài toán.", sz=13, c=INK, line=1.08)

card(s, 6.75, 1.30, 6.05, 2.72)
simple(s, 7.01, 1.50, 5.53, 0.32, "BỐN THUẬT TOÁN, CÙNG MỘT MÔI TRƯỜNG",
       sz=12.5, b=True, c=ORANGE, wrap=False)
simple(s, 7.01, 1.88, 5.53, 0.30, "Nhóm giải được câu đố Capture The Flag",
       sz=13, c=MUTED, wrap=False)
textbox(s, 7.01, 2.26, 5.53, 0.80, [
    {"runs": [("0,800", {"sz": 40, "b": True, "c": NAVY}),
              ("     –     ", {"sz": 26, "c": MUTED}),
              ("0,840", {"sz": 40, "b": True, "c": NAVY})], "align": "c", "line": 1.0}])
simple(s, 7.01, 3.14, 5.53, 0.72,
       "PPO, MA-POCA, MAPPO và DQN nằm gọn trong dải rộng 0,040 — nhỏ hơn tới mười sáu lần.",
       sz=13, c=INK, line=1.08)

round_rect(s, 0.55, 4.20, 12.25, 0.92, fill=CARD2, adj=0.16)
textbox(s, 0.75, 4.20, 11.85, 0.92, [
    {"runs": [("0,663", {"sz": 24, "b": True, "c": ORANGE}),
              ("   chênh lệch do môi trường        so với        ", {"sz": 15.5, "c": INK}),
              ("0,040", {"sz": 24, "b": True, "c": NAVY}),
              ("   chênh lệch do thuật toán", {"sz": 15.5, "c": INK})],
     "align": "c", "line": 1.0}], anchor="c")

card(s, 0.55, 5.32, 12.25, 1.56)
textbox(s, 0.86, 5.32, 11.63, 1.56, [
    {"runs": [("Hàm ý cho người làm game:  ", {"sz": 15, "b": True, "c": ORANGE}),
              ("cơ chế game trước, thuật toán sau.", {"sz": 15})],
     "line": 1.10, "space_after": 7},
    {"runs": [("Ba phân tách vững:  ", {"sz": 13.5, "b": True, "c": NAVY}),
              ("SAC thất bại ở Capture The Flag  ·  DQN cuối bảng ở Cross The Road  ·  nhóm dẫn "
               "đầu tách khỏi nhóm sau ở Cross The Road.", {"sz": 13.5, "c": MUTED})],
     "line": 1.08}], anchor="c")
notes(s, """Đây là luận điểm trung tâm của đồ án, và em xin dừng lại hơi lâu ở slide này.

Nhìn sang trái. Cùng một thuật toán là SAC, cùng phần cứng, cùng đường ống đo, cùng người cấu
hình. Trên Cross The Road nó dẫn đầu với không phẩy tám bốn sáu. Trên Capture The Flag nó sụp
xuống không phẩy một tám ba. Chênh lệch không phẩy sáu sáu ba, và chênh lệch đó sinh ra hoàn toàn
từ đặc tính bài toán mà thiết kế môi trường quyết định.

Nhìn sang phải. Bốn thuật toán khác nguyên lý — PPO là học chính sách theo lô, MA-POCA và MAPPO có
phê bình tập trung, DQN học hàm giá trị theo hướng off-policy với bộ nhớ phát lại — cùng giải được Capture The Flag và nằm gọn
trong dải rộng đúng không phẩy không bốn.

Đặt cạnh nhau: chênh lệch do môi trường lớn gấp khoảng mười sáu lần chênh lệch do thuật toán. Đây
là bằng chứng định lượng cho điều em nêu ở slide bối cảnh.

Hàm ý thực tiễn cho người làm game là hãy lo cơ chế và đặc tả trước, chọn thuật toán sau. Chọn sai
thuật toán cho một cơ chế game gây thiệt hại lớn hơn nhiều so với cái lợi của việc chọn được thuật
toán tốt nhất trong nhóm khả dụng.

Em cũng xin nói rõ giới hạn của kết luận này. Chỉ có ba phân tách là thực sự vững: SAC thất bại ở
Capture The Flag, DQN cuối bảng ở Cross The Road, và nhóm dẫn đầu tách khỏi nhóm sau ở Cross The
Road. Còn thứ tự trong nội bộ nhóm dẫn đầu, ví dụ MAPPO không phẩy tám bốn so với MA-POCA không
phẩy tám ba tám, thì chênh lệch nhỏ hơn dao động của đường cong nên em không coi đó là kết luận.""")

# ============================================ 17. SẢN PHẨM TRIỂN KHAI & DEMO ===
s = new_slide_t("Sản phẩm triển khai và ứng dụng demo")
picture_plate(s, 0.55, 1.26, 6.90, 4.28, os.path.join(IMG, "app_menu_main.png"), pad=0.12)
caption(s, 0.55, 5.62, 6.90,
        "Tab CHƠI: khung xem trước 3D, lưới chọn chế độ và danh mục mô hình nạp từ Resources.")

card(s, 0.55, 6.02, 6.90, 0.86)
simple(s, 0.82, 6.02, 6.36, 0.86,
       "Quan sát hành vi thật phát hiện lỗi đặc tả mà đường cong phần thưởng che giấu.",
       sz=13, b=True, c=NAVY, anchor="c", line=1.08)

_prod = [("Gói Unity Package Manager", "3 môi trường · 15 cấu hình · công cụ Editor"),
         ("Ứng dụng demo nạp ONNX", "15 mô hình chọn trực tiếp từ menu, không dựng lại"),
         ("Ba chế độ chơi", "người chơi · agent tự chơi · người và agent"),
         ("Hai tab số liệu", "THÔNG SỐ và SO SÁNH đọc thẳng kết quả huấn luyện")]
for i, (h, b) in enumerate(_prod):
    headline(s, 7.80, 1.26 + i * 1.44, 5.00, 1.30, h, b, hsz=15, bsz=12.5)
notes(s, """Sản phẩm bàn giao của đồ án gồm hai phần.

Phần một là gói Unity Package Manager. Người dùng cài gói này vào một dự án Unity trống là có ngay
ba môi trường, mười lăm tệp cấu hình đã chạy thật, bộ công cụ Editor và một trình hướng dẫn cài
đặt bảy bước. Mục tiêu là để người khác dựng lại được thí nghiệm mà không phải làm lại từ đầu.

Phần hai là ứng dụng demo trong ảnh. Ứng dụng nạp trực tiếp mười lăm mô hình ONNX đã huấn luyện
tại thời gian chạy, chọn ngay từ menu, không cần dựng lại bản build. Có ba chế độ chơi: người chơi
tự điều khiển, agent tự chơi, và người chơi cùng agent trong một phiên. Ngoài ra có hai tab số
liệu đọc thẳng kết quả huấn luyện, nên số trên slide và số trong demo luôn khớp nhau vì cùng một
nguồn.

Điều em muốn nhấn mạnh là giá trị của việc xem hành vi thật, ghi ở dòng dưới cùng bên trái. Có
những lỗi đặc tả mà đường cong phần thưởng hoàn toàn che giấu: phần thưởng vẫn tăng, đường cong
vẫn đẹp, nhưng nhìn tác nhân chạy thì thấy ngay nó đang làm một việc khác với điều mình muốn. Ứng
dụng demo chính là công cụ kiểm tra tiêu chí "không có lối tắt" mà em nêu ở phần nghiệm thu.

Nếu hội đồng cho phép, sau phần trình bày em xin demo trực tiếp một môi trường.""")

# ============================== 18. KẾT LUẬN, HẠN CHẾ, HƯỚNG PHÁT TRIỂN ========
s = new_slide_t("Kết luận, hạn chế và hướng phát triển")

_cols = [
    ("KẾT QUẢ ĐẠT ĐƯỢC", NAVY, "✓", [
        ("Ba môi trường đã qua kiểm chứng", "mỗi môi trường chịu 5 thuật toán khác nguyên lý"),
        ("Quy trình thiết kế và nghiệm thu", "4 nguyên tắc phần thưởng, 3 tiêu chí nghiệm thu"),
        ("Sản phẩm hoàn chỉnh", "gói UPM, trainer mở rộng, ứng dụng demo"),
    ]),
    ("HẠN CHẾ", RED, "!", [
        ("Mỗi cấu hình chỉ chạy một seed", "chưa tách được bản chất thuật toán khỏi may rủi"),
        ("Hai môi trường thiếu độ phân giải", "dải kết quả hẹp hơn dao động đường cong"),
        ("Một ô nằm ngoài phép so sánh", "DQN trên Football Table chạy không tự chơi"),
    ]),
    ("HƯỚNG PHÁT TRIỂN", GREEN, "→", [
        ("Chạy lại 15 tổ hợp trên nhiều seed", "việc cần làm trước tiên"),
        ("Làm hai môi trường khó lên", "thêm bài trong chương trình học, buộc phối hợp chặt hơn"),
        ("Kiểm chứng hai hiện tượng bỏ ngỏ", "SAC sụp ở CTF; MAPPO vượt PPO ở môi trường đơn"),
    ]),
]
for i, (h, c, mark, items) in enumerate(_cols):
    x = 0.55 + i * 4.15
    card(s, x, 1.26, 3.95, 4.66)
    simple(s, x + 0.26, 1.46, 3.43, 0.34, h, sz=15, b=True, c=c, wrap=False)
    _iy = 1.96
    for ih, ib in items:
        oval(s, x + 0.26, _iy + 0.04, 0.30, fill=c, label=mark, sz=11)
        textbox(s, x + 0.66, _iy, 3.03, 1.24, [
            {"runs": [(ih, {"sz": 13.5, "b": True, "c": NAVY})], "line": 1.02, "space_after": 4},
            {"runs": [(ib, {"sz": 12, "c": MUTED})], "line": 1.06}])
        _iy += 1.30

card(s, 0.55, 6.06, 12.25, 0.82)
simple(s, 0.82, 6.06, 11.71, 0.82,
       "Đóng góp chính không phải một thuật toán mới, mà là ba môi trường kèm bằng chứng nghiệm "
       "thu và quy trình dựng lại chúng.",
       sz=14.5, b=True, c=NAVY, align="c", anchor="c", wrap=False)
notes(s, """Em xin tổng kết.

Về kết quả đạt được, đồ án bàn giao ba môi trường đã qua kiểm chứng, mỗi môi trường chịu được năm
thuật toán khác nguyên lý; một quy trình thiết kế và nghiệm thu gồm bốn nguyên tắc phần thưởng và
ba tiêu chí nghiệm thu; cùng sản phẩm hoàn chỉnh gồm gói UPM, hai trainer mở rộng và ứng dụng
demo.

Về hạn chế, em xin nêu trung thực và xếp theo mức độ nghiêm trọng. Hạn chế lớn nhất là mỗi cấu
hình mới chỉ chạy một seed. Điều đó có nghĩa là những chênh lệch sát nhau — ví dụ MAPPO so với
MA-POCA ở Capture The Flag — chưa đủ cơ sở thống kê để khẳng định. Nếu chỉ được sửa một điều trong
đồ án này, em sẽ sửa điều đó. Hạn chế thứ hai là hai môi trường Football Table và Capture The Flag
có dải kết quả hẹp hơn dao động của chính đường cong, nên chưa đủ độ phân giải để xếp hạng. Hạn
chế thứ ba là ô DQN trên Football Table chạy trong điều kiện khác, nên nằm ngoài phép so sánh.

Về hướng phát triển, việc cần làm trước tiên là chạy lại toàn bộ mười lăm tổ hợp trên nhiều seed,
vì nó quyết định mọi kết luận về chênh lệch nhỏ. Tiếp theo là làm hai môi trường thiếu độ phân
giải khó lên, bằng cách thêm bài trong chương trình học hoặc buộc phối hợp chặt hơn, thay vì kéo
dài thời lượng huấn luyện. Cuối cùng là kiểm chứng hai hiện tượng còn bỏ ngỏ mà em đã nêu.

Câu ở dưới cùng là điều em muốn hội đồng ghi nhận: đóng góp chính của đồ án không phải một thuật
toán mới, mà là ba môi trường kèm bằng chứng nghiệm thu và quy trình để dựng lại chúng.""")

# ================================================================ 19. KẾT THÚC =
s = blank_title_slide()
round_rect(s, 1.80, 2.62, 9.73, 1.10, fill=None, line=AMBER, lw=1.25, adj=0.14)
simple(s, 2.05, 2.62, 9.23, 1.10,
       "Muốn học tăng cường trong game đáng tin cậy, cần đầu tư vào đặc tả và kiểm chứng "
       "môi trường trước khi tối ưu thuật toán.",
       sz=17, b=True, c=WHITE, align="c", anchor="c", line=1.16)
simple(s, 1.50, 4.06, 10.33, 0.62, "EM XIN CHÂN THÀNH CẢM ƠN",
       sz=32, b=True, c=WHITE, align="c", wrap=False)
simple(s, 2.00, 4.76, 9.33, 0.36, "sự lắng nghe của quý thầy cô trong hội đồng",
       sz=16, c=SKY, align="c", wrap=False)
rect(s, 5.67, 5.34, 2.00, 0.02, fill=AMBER)
simple(s, 2.00, 5.58, 9.33, 0.36, "Kính mong nhận được góp ý của thầy cô để đồ án được hoàn thiện hơn.",
       sz=14, c=SKY, align="c", wrap=False)
simple(s, 2.00, 6.16, 9.33, 0.36, "Văn Thị Minh Huyền  ·  22010329  ·  K16  ·  Khoa học máy tính",
       sz=14, b=True, c=WHITE, align="c", wrap=False)
notes(s, """Phần trình bày của em đến đây là hết.

Thông điệp em muốn để lại là câu trong khung: muốn học tăng cường trong game đáng tin cậy thì phải
đầu tư vào đặc tả và kiểm chứng môi trường trước khi nghĩ tới việc tối ưu thuật toán. Toàn bộ số
liệu trong bài đều hướng về kết luận đó.

Em xin chân thành cảm ơn quý thầy cô đã lắng nghe, và rất mong nhận được góp ý để đồ án được hoàn
thiện hơn. Em xin sẵn sàng trả lời câu hỏi của hội đồng.""")

# ------------------------------------------------------------------------ lưu
prs.save(OUT)
print("Đã ghi:", OUT, "-", len(prs.slides._sldIdLst), "slide")
_words = sum(len(sl.notes_slide.notes_text_frame.text.split()) for sl in prs.slides)
print(f"Kịch bản nói: {_words} tiếng  (~{_words / 180:.0f} phút ở nhịp 180 tiếng/phút)")
if OVERFLOW:
    print("\n[CẢNH BÁO TRÀN CHỮ]")
    for w in OVERFLOW:
        print(" -", w)
else:
    print("Không phát hiện tràn chữ theo ước lượng.")
