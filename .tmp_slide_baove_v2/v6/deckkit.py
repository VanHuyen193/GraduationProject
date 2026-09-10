# -*- coding: utf-8 -*-
"""Helpers dựng slide theo đúng hệ thiết kế của template Phenikaa trong bản v4."""
import os
from lxml import etree
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import ImageFont

# ---------------------------------------------------------------- design tokens
NAVY   = RGBColor(0x22, 0x37, 0x71)
ORANGE = RGBColor(0xF2, 0x65, 0x22)
CARD   = RGBColor(0xF1, 0xF4, 0xFA)
CARD2  = RGBColor(0xE8, 0xEC, 0xF6)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
INK    = RGBColor(0x2B, 0x2B, 0x2B)
MUTED  = RGBColor(0x66, 0x6E, 0x7E)
LINE   = RGBColor(0xD5, 0xDC, 0xEA)
ICE    = RGBColor(0xE8, 0xEC, 0xF3)

FONT = "Calibri"
SLIDE_W, SLIDE_H = 13.3333, 7.5

# ------------------------------------------------------------------ text metrics
_FONTS = {}
_PATHS = {
    (False, False): r"C:\Windows\Fonts\calibri.ttf",
    (True,  False): r"C:\Windows\Fonts\calibrib.ttf",
    (False, True):  r"C:\Windows\Fonts\calibrii.ttf",
    (True,  True):  r"C:\Windows\Fonts\calibriz.ttf",
}


def _font(pt, bold=False, italic=False):
    key = (round(pt * 2), bold, italic)
    if key not in _FONTS:
        path = _PATHS[(bold, italic)]
        if not os.path.exists(path):
            path = _PATHS[(False, False)]
        _FONTS[key] = ImageFont.truetype(path, max(4, int(round(pt * 4))))
    return _FONTS[key]


def text_w(s, pt, bold=False, italic=False):
    """Bề rộng chuỗi, tính bằng inch."""
    return _font(pt, bold, italic).getlength(s) / 4.0 / 72.0


def wrap_count(s, pt, width_in, bold=False, italic=False):
    """Số dòng khi bọc chữ trong khung rộng width_in inch."""
    lines, cur = 0, ""
    for word in s.split():
        trial = word if not cur else cur + " " + word
        if text_w(trial, pt, bold, italic) <= width_in or not cur:
            cur = trial
        else:
            lines += 1
            cur = word
    return lines + (1 if cur else 0)


# ------------------------------------------------------------------- primitives
def _clear_style(shape):
    shape.shadow.inherit = False


def round_rect(slide, x, y, w, h, fill=CARD, line=None, lw=1.0, adj=0.07):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = adj
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    _clear_style(sh)
    sh.text_frame.text = ""
    return sh


def rect(slide, x, y, w, h, fill=CARD, line=None, lw=1.0):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    _clear_style(sh)
    sh.text_frame.text = ""
    return sh


def oval(slide, x, y, d, fill=NAVY, label=None, sz=11.5, color=WHITE, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                                Inches(d), Inches(d))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.0)
    _clear_style(sh)
    tf = sh.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label or ""
    r.font.size = Pt(sz)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = color
    return sh


def arrow(slide, x, y, w, h, shape=MSO_SHAPE.RIGHT_ARROW, fill=None):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill or RGBColor(0xB4, 0xBE, 0xD6)
    sh.line.fill.background()
    _clear_style(sh)
    sh.text_frame.text = ""
    return sh


def textbox(slide, x, y, w, h, paras, anchor="t", margins=(0, 0, 0, 0), wrap=True):
    """paras: list of dict(runs=[(text, opts)], align, space_before, space_after, line)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [
        Inches(v) for v in margins]
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "c": MSO_ANCHOR.MIDDLE,
                          "b": MSO_ANCHOR.BOTTOM}[anchor]
    for i, spec in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                       "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}[spec.get("align", "l")]
        if spec.get("space_before") is not None:
            p.space_before = Pt(spec["space_before"])
        else:
            p.space_before = Pt(0)
        p.space_after = Pt(spec.get("space_after", 0))
        p.line_spacing = spec.get("line", 1.0)
        runs = spec.get("runs") or [(spec.get("t", ""), {})]
        for text, opt in runs:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(opt.get("sz", 15))
            r.font.bold = opt.get("b", False)
            r.font.italic = opt.get("i", False)
            r.font.name = FONT
            r.font.color.rgb = opt.get("c", INK)
            if opt.get("sub"):
                r.font._rPr.set("baseline", "-25000")
    return tb


def dash_bullets(slide, x, y, w, h, items, sz=13.5, c=INK, bu=None,
                 line=1.08, space_after=6, char="–", marL=0.24):
    """Danh sách gạch đầu dòng có thụt lề treo, để dòng tràn không tụt về lề trái."""
    bu = bu or ORANGE
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line
        p.space_before = Pt(0)
        p.space_after = Pt(space_after if i < len(items) - 1 else 0)
        r = p.add_run()
        r.text = item
        r.font.size = Pt(sz)
        r.font.name = FONT
        r.font.color.rgb = c
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(marL))))
        pPr.set("indent", str(-int(Inches(marL))))
        buClr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(buClr, qn("a:srgbClr")).set("val", str(bu))
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", FONT)
        etree.SubElement(pPr, qn("a:buChar")).set("char", char)
    return tb


def simple(slide, x, y, w, h, text, sz=15, b=False, i=False, c=INK,
           align="l", anchor="t", line=1.0, margins=(0, 0, 0, 0), wrap=True):
    return textbox(slide, x, y, w, h,
                   [{"runs": [(text, {"sz": sz, "b": b, "i": i, "c": c})],
                     "align": align, "line": line}],
                   anchor=anchor, margins=margins, wrap=wrap)


# ------------------------------------------------------------- composite pieces
def title(slide, text, sz=26, w=9.90):
    tb = simple(slide, 0.42, 0.20, w, 0.72, text, sz=sz, b=True, c=NAVY,
                anchor="c", wrap=False)
    return tb


def chip(slide, x, y, text, sz=15.5, color=NAVY, h=0.42, w=None):
    """Nhãn mục: hộp bo góc viền màu, chữ đậm căn giữa. Bề rộng đo theo font thật."""
    if w is None:
        w = round(text_w(text, sz, bold=True) + 0.46, 2)
    box = round_rect(slide, x, y, w, h, fill=WHITE, line=color, lw=1.25, adj=0.16)
    simple(slide, x, y, w, h, text, sz=sz, b=True, c=color, align="c",
           anchor="c", wrap=False)
    return box, w


def card(slide, x, y, w, h, fill=CARD, adj=0.07):
    return round_rect(slide, x, y, w, h, fill=fill, adj=adj)


def callout(slide, x, y, w, h, color=ORANGE, adj=0.04):
    return round_rect(slide, x, y, w, h, fill=WHITE, line=color, lw=1.0, adj=adj)


def head_body(slide, x, y, w, head, body, hsz=16.5, bsz=15, gap=5,
              hcolor=NAVY, bcolor=INK, line=1.06, h=None, anchor="t"):
    """Tiêu đề đậm + đoạn mô tả trong một hộp chữ."""
    paras = [{"runs": [(head, {"sz": hsz, "b": True, "c": hcolor})],
              "line": 1.0, "space_after": gap}]
    if body:
        paras.append({"runs": [(body, {"sz": bsz, "c": bcolor})], "line": line})
    est = (hsz * 1.30 * wrap_count(head, hsz, w, bold=True) + gap
           + (bsz * 1.30 * line * wrap_count(body, bsz, w) if body else 0)) / 72.0
    return textbox(slide, x, y, w, h or (est + 0.06), paras, anchor=anchor), est


def card_text(slide, cx, cy, cw, chh, head, body, hsz=16, bsz=14, gap=5,
              padx=0.26, pady=0.18, hcolor=NAVY, line=1.08):
    """Thẻ nội dung với chữ căn giữa theo chiều dọc, tránh khoảng trống dưới đáy."""
    card(slide, cx, cy, cw, chh)
    _, est = head_body(slide, cx + padx, cy + pady, cw - 2 * padx, head, body,
                       hsz=hsz, bsz=bsz, gap=gap, hcolor=hcolor, line=line,
                       h=chh - 2 * pady, anchor="c")
    return est


def spec_row(slide, x, y, w, h, label, value, body, lsz=14, vsz=13, bsz=13.5):
    """Hàng thông số: nhãn đậm + giá trị nhấn cùng dòng, mô tả bên dưới."""
    card(slide, x, y, w, h)
    tw = w - 0.48
    textbox(slide, x + 0.24, y + 0.14, tw, h - 0.28, [
        {"runs": [(label + "   ", {"sz": lsz, "b": True, "c": NAVY}),
                  (value, {"sz": vsz, "b": True, "c": ORANGE})],
         "line": 1.0, "space_after": 4},
        {"runs": [(body, {"sz": bsz, "c": INK})], "line": 1.08},
    ], anchor="c")


def stat_tile(slide, x, y, w, h, number, label, nsz=30, lsz=12.5):
    card(slide, x, y, w, h)
    simple(slide, x, y + 0.10, w, h * 0.55, number, sz=nsz, b=True, c=NAVY,
           align="c", wrap=False)
    simple(slide, x, y + h - 0.46, w, 0.40, label, sz=lsz, c=MUTED, align="c")


def caption(slide, x, y, w, text, sz=12, c=MUTED, align="l"):
    return simple(slide, x, y, w, 0.30, text, sz=sz, i=True, c=c, align=align)


def picture_plate(slide, x, y, w, h, img, pad=0.18):
    """Ảnh đặt trong nền trắng bo góc, giữ tỉ lệ và căn giữa."""
    round_rect(slide, x, y, w, h, fill=WHITE, adj=0.06)
    from PIL import Image
    iw, ih = Image.open(img).size
    ar = iw / ih
    aw, ah = w - 2 * pad, h - 2 * pad
    if aw / ah > ar:
        ph = ah
        pw = ah * ar
    else:
        pw = aw
        ph = aw / ar
    slide.shapes.add_picture(img, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2),
                             Inches(pw), Inches(ph))
    return pw, ph


def table(slide, x, y, w, h, data, col_w=None, header_fill=NAVY,
          hsz=13, bsz=13, row_h=0.34, band=CARD, first_col_bold=True,
          first_col_color=None, align_first="c"):
    rows, cols = len(data), len(data[0])
    gf = slide.shapes.add_table(rows, cols, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    if col_w:
        total = sum(col_w)
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Emu(int(Inches(w) * cw / total))
    for r in range(rows):
        tbl.rows[r].height = Inches(row_h)
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = header_fill
            else:
                cell.fill.fore_color.rgb = WHITE if r % 2 else band
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if (r == 0 or c > 0 or align_first == "c") \
                else PP_ALIGN.LEFT
            run = p.add_run()
            run.text = str(data[r][c])
            run.font.name = FONT
            run.font.size = Pt(hsz if r == 0 else bsz)
            run.font.bold = (r == 0) or (c == 0 and first_col_bold)
            if r == 0:
                run.font.color.rgb = WHITE
            elif c == 0 and first_col_color:
                run.font.color.rgb = first_col_color(data[r][0])
            else:
                run.font.color.rgb = INK
    return tbl


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text
